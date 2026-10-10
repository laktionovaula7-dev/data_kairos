# -*- coding: utf-8 -*-
"""
Регенератор дашборда Кайрос.
Читает актуальные Excel из папки, считает всё сам (ABC, оборачиваемость, год-к-году),
вставляет данные в шаблон (Kairos_dashboard_final.html) и пересобирает дашборд.

Запуск:  python rebuild.py
Обновление: положить свежие выгрузки в папку с теми же именами -> запустить скрипт.

Охват: ВСЯ компания, все бренды (HeadRock/KRONbuild/ENKI), все регионы/страны, 2024-2026.
"""
import sys, re, json, zipfile, pickle, time, os
import xml.etree.ElementTree as ET
from collections import defaultdict
from pathlib import Path
from datetime import date, datetime

try: sys.stdout.reconfigure(encoding="utf-8")
except Exception: pass

ROOT   = Path(__file__).resolve().parent
DATA   = ROOT/"Продажи и остатки"                      # свежие выгрузки кладём сюда (fallback — корень)
def _pick(*cands):
    """Первый существующий путь из кандидатов (свежая папка приоритетнее корня)."""
    for p in cands:
        if p and Path(p).exists(): return Path(p)
    return Path(cands[0])
# Продажи: полная история (2024-2026) — из неё берём только 2024-2025; свежий файл — 2026 до актуальной даты
F_SALES= _pick(DATA/"Продажи по бизнес регионам 24г-26г.xlsx", ROOT/"Продажи по бизнес регионам 24г-26г.xlsx")
F_SALES_CUR=_pick(DATA/"продажи 09.10.xlsx", DATA/"продажи 08.10.xlsx", DATA/"Продажи по Бизнес регионам 01.01.26-05.10.26.xlsx")   # накопленный 2026 (полный период 01.01–<дата>)
F_STOCK= _pick(DATA/"остатки по сериям 09.10.xlsx", DATA/"Остатки 08.10.xlsx", DATA/"Остатки и доступность товаров (по сериям) 05.10.xlsx",
               DATA/"Остатки и доступность по сериям.xlsx", ROOT/"Остатки и доступность по сериям.xlsx")
F_EXPIRY=ROOT/"Отчет по товарам на складах с окончанием срока годности.xlsx"
F_PRICE= ROOT/"Прайс NEW HR+KRON+ENKI МСК-ВЛ (03 Сентября 2026г) (1).xlsx"
CRM_DIR= ROOT/"Регионы"                                # выгрузки CRM по округам (ЮФО.xls, ПФО.xls, …, ЦФО.files)
F_PLAN = _pick(DATA/"План_25_26.xlsx", ROOT/"План_25_26.xlsx")
# Слежение / в пути: заказы поставщику (Китай) со статусами «в производстве / в пути» и датами поступления
F_TRACK= _pick(DATA/"в пути 09.10.xlsx", DATA/"в пути и в производстве.xlsx",
               DATA/"Остатки и доступность товаров (слежение) с датами выхода, прихода, поступления..xlsx",
               DATA/"Остатки и доступность товаров (слежение) с датами выхода, прихода, поступления.xlsx")

def _find_client_inn():
    if CRM_DIR.exists():
        for p in CRM_DIR.glob("*лиент*ИНН*.xlsx"): return p
    return None
TEMPLATE_SRC = ROOT/"Kairos_dashboard_final.html"      # источник вёрстки (из него берём слот данных)
OUT    = ROOT/"Kairos_dashboard_rebuilt.html"          # результат (отдельный файл до проверки)
CACHE  = ROOT/"work"/"_sales_rows2.pkl"                # кэш сырых строк продаж (полная история)
CACHE_CUR=ROOT/"work"/"_sales_rows_cur.pkl"            # кэш сырых строк продаж (свежий 2026)

CUR_YEAR, PREV_YEAR = 2026, 2025
TURN_WINDOW_DAYS = 90          # окно для оборачиваемости (DOS)

# Правило: исключаем ВСЕ компании с нулевыми (и отрицательными) продажами.
# Служебные записи (сотрудники компании) — исключаем из клиентской аналитики.
EXCLUDE_CLIENTS_RAW = ['Евгений Тихонов','Бурилова Валерия Викторовна']
# Ручная привязка нераспределённых клиентов: nrm(имя) -> (регион, филиал, субъект, город).
# Менеджер проставляется автоматически по региону (OKRUG_MGR): ЦФО→Сидоров, ДВФО→Федотов.
MANUAL_CLIENT_GEO = {
    'дементиенко роза гиясовна':   ('Дальневосточный ФО','Владивосток','Амурская область','Тында'),
    'фастрост ооо':                ('Центральный ФО','Москва','Ярославская область','Ярославль'),
    'гробман евгений борисович':   ('Дальневосточный ФО','Владивосток','Забайкальский край','Чита'),
    'богатыренко эдуард сергеевич':('Центральный ФО','Москва','Калужская область','Кондрово'),
    'слк ооо':                     ('Дальневосточный ФО','Владивосток','Республика Бурятия','Улан-Удэ'),
}
MN=['Янв','Фев','Мар','Апр','Май','Июн','Июл','Авг','Сен','Окт','Ноя','Дек']

def nrm(s):
    s=str(s or '').lower().strip().replace('ё','е')
    s=re.sub(r'[«»"\'`(),]',' ',s); s=s.replace('х','x'); s=re.sub(r'\s+',' ',s)
    return s.strip()

# ---------------- ПРАЙС: имя/артикул -> бренд, категория ----------------
def load_price():
    import openpyxl
    if not F_PRICE.exists(): print("  [!] нет прайса"); return {},{}
    wb=openpyxl.load_workbook(F_PRICE,data_only=True,read_only=True); rows=[]
    art_price={}
    def _pr(x):
        try: return float(x)
        except (TypeError,ValueError): return 0.0
    if 'HeadRock' in wb.sheetnames:
        ws=wb['HeadRock']; cat=None
        for r in ws.iter_rows(min_row=9,values_only=True):
            art=r[1] if len(r)>1 else None; name=r[2] if len(r)>2 else None
            if art and (name is None or str(name).strip()==''): cat=str(art).strip(); continue
            if art and name:
                a=str(art).strip(); rows.append((a,str(name).strip(),'HeadRock',cat))
                if len(r)>6: art_price[a]=_pr(r[6])          # колонка G = цена
    if 'KRONbuild' in wb.sheetnames:
        ws=wb['KRONbuild']; cat=None
        for r in ws.iter_rows(min_row=10,values_only=True):
            art=r[0] if len(r)>0 else None; name=r[2] if len(r)>2 else None
            if (art is None or str(art).strip()=='') and name and str(name).strip(): cat=str(name).strip(); continue
            if art and name:
                brand='ENKI' if (cat and 'ENKI' in cat.upper()) else 'KRONbuild'
                a=str(art).strip(); rows.append((a,str(name).strip(),brand,cat))
                if len(r)>6: art_price[a]=_pr(r[6])
    wb.close()
    price={nrm(n):(a,n,b,c) for a,n,b,c in rows}
    price['__art_price__']=art_price                        # цена по артикулу
    print(f"  прайс: {len(rows)} SKU, цен по артикулу: {len(art_price)}")
    return price

def load_geo():
    f=ROOT/"geo_ru.json"
    if f.exists():
        try:
            g=json.loads(f.read_text(encoding='utf-8')); g.pop('_note',None)
            n=sum(len(cs) for sub in g.values() for cs in sub.values())
            print(f"  гео-справочник: округов {len(g)}, городов {n}")
            return g
        except Exception as e: print("  [!] geo_ru.json:",e)
    return {}

def load_rumap():
    f=ROOT/"geo_rumap.json"
    if f.exists():
        try:
            m=json.loads(f.read_text(encoding='utf-8'))
            print(f"  карта РФ: регионов {len(m.get('regions',[]))}")
            return m
        except Exception as e: print("  [!] geo_rumap.json:",e)
    return {}

def load_kontur():
    """ИНН-справочник (имя из 1С -> ИНН) + кэш обогащения Контур.Фокус (ИНН -> данные)."""
    cinn={}
    f=ROOT/"work"/"_client_inn.xlsx"
    if f.exists():
        try:
            import openpyxl
            wb=openpyxl.load_workbook(f,read_only=True,data_only=True)
            for r in wb.active.iter_rows(min_row=2,values_only=True):
                nm=r[0] if len(r)>0 else None; inn=r[2] if len(r)>2 else None
                if nm and inn: cinn[nrm(str(nm))]=str(inn).strip()
            wb.close()
        except Exception as e: print("  [!] _client_inn.xlsx:",e)
    kont={}
    kf=ROOT/"work"/"_kontur_cache.json"
    if kf.exists():
        try: kont=json.loads(kf.read_text(encoding='utf-8'))
        except Exception as e: print("  [!] _kontur_cache.json:",e)
    print(f"  Контур: ИНН-справочник {len(cinn)} имён · обогащено компаний {len(kont)}")
    return cinn, kont

def brand_of(name, price):
    m=price.get(nrm(name))
    if m: return m[2]
    u=name.lower()
    if 'kronbuild' in u or 'кронбилд' in u: return 'KRONbuild'
    if 'enki' in u or 'энки' in u: return 'ENKI'
    if 'headrock' in u or 'хедрок' in u: return 'HeadRock'
    return 'HeadRock'   # инструмент по умолчанию — HeadRock

POSM_KW=('каталог','буклет','брошюр','плакат','воблер','листовк','ценник','наклейк','пакет','стенд','полиграф','посм','pos-','pos ','сумка','ролл-ап','ролап','штендер','флаер')
def is_posm(name):
    n=(name or '').lower()
    return any(k in n for k in POSM_KW)

def cat_art(name, price):
    m=price.get(nrm(name)); return (m[3],m[0]) if m else (None,None)

def group_of(cat, brand):
    """Товарная группа (верхний уровень) по категории прайса. Подгруппа = сама категория прайса.
    HeadRock — 12 групп; KRON/ENKI — 6 групп; пистолеты (подбренд HeadRock) → «Отделка и нанесение»."""
    s=(cat or '').lower()
    if brand in ('KRONbuild','ENKI'):
        if s.startswith('пистолет'): return 'Отделка и нанесение'
        if 'уплотнит' in s: return 'Уплотнители'
        if 'гермет' in s: return 'Герметики'
        if 'пен' in s: return 'Монтажные пены'
        if 'жидкие гвозди' in s or 'кле' in s: return 'Клеи'
        if 'эмал' in s or 'грунт' in s or 'краск' in s or 'лак' in s: return 'ЛКМ'
        if 'шуруп' in s or 'дюбел' in s or 'анкер' in s or 'саморез' in s: return 'Крепёж'
        return 'Прочее'
    # HeadRock
    if 'коронк' in s or 'кольцевых пил' in s or 'балеринк' in s: return 'Коронки'
    if 'пика sds' in s or 'зубило' in s: return 'Пики и зубила (SDS)'
    if 'бур по бетону' in s or 'сверло' in s or 'шнеков' in s or 'перов' in s or 'удлинитель для пер' in s or 'конфирмат' in s or 'борфрез' in s: return 'Буры и свёрла'
    if 'полотна' in s or 'цепь' in s or 'шина' in s or 'ножовк' in s or 'напильник' in s: return 'Полотна, цепи, ножовки'
    if 'войлоч' in s or 'полировал' in s: return 'Абразив и полировка'
    if 'диск' in s: return 'Диски'
    if 'щетк' in s or 'щётк' in s or 'зачистн' in s or 'шарошек' in s: return 'Щётки и зачистка'
    if 'шлифовал' in s or 'лепестков' in s or 'черепашк' in s or 'губки абразив' in s or 'насадка' in s or 'чаша алмазн' in s or 'быстрозажимная гайка' in s: return 'Абразив и полировка'
    if 'биты' in s or 'ответок' in s or 'редуктор' in s or 'звёздоч' in s or 'шестигранник' in s or 'трещотк' in s or 'патрон' in s: return 'Биты и головки'
    if 'шпател' in s or 'гладилк' in s or 'ручка телескоп' in s or 'скребок' in s or 'пистолет' in s or 'маркер' in s or 'карандаш' in s: return 'Отделка и нанесение'
    if 'рулетк' in s or 'уровни' in s or 'штангенциркул' in s or 'отбивочн' in s: return 'Измерение и разметка'
    if 'стяжки кабельн' in s or 'лески' in s or 'стержни клеев' in s: return 'Расходники'
    if any(w in s for w in ['струбцин','стяжные ремн','пассатижи','тонкогубц','бокорез','клещи','ключ','нож','лезви','степлер','скобы','горелк','крюк']): return 'Ручной инструмент'
    return 'Прочее'

def group_cat(cat, brand, name=None):
    """Укрупняем категории KRON/ENKI в верхние группы (по прайс-категории, а если её нет — по названию); HeadRock — по прайсу как есть."""
    def _infer(t):
        if not t: return None
        if 'пена' in t: return 'Монтажные пены'
        if 'герметик' in t: return 'Герметики'
        if 'клей' in t or 'жидкие гвозд' in t: return 'Клей'
        if 'краск' in t or 'эмал' in t or 'грунт' in t or 'аэрозол' in t: return 'Аэрозольные краски'
        if 'пистолет' in t: return 'Пистолеты и оснастка'
        return None
    if brand in ('KRONbuild','ENKI'):
        g=_infer((cat or '').lower()) or _infer((name or '').lower())
        if g: return g
    return cat

# ---------------- ПРОДАЖИ: быстрый загрузчик (значения read_only + уровни из XML) --------
def load_sales_rows(path=None, cache=None):
    path=Path(path) if path else F_SALES
    cache=Path(cache) if cache else CACHE
    if cache.exists() and cache.stat().st_mtime>=path.stat().st_mtime:
        print(f"  продажи ({path.name}): из кэша"); return pickle.loads(cache.read_bytes())
    import openpyxl
    t0=time.time()
    ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    levels={}
    with zipfile.ZipFile(path).open('xl/worksheets/sheet1.xml') as fp:
        for ev,el in ET.iterparse(fp,events=('end',)):
            if el.tag==ns+'row':
                r=int(el.get('r')); ol=el.get('outlineLevel'); levels[r]=int(ol) if ol else 0
                el.clear()
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True); ws=wb['Лист_1']
    rows=[]; r=0
    for v in ws.iter_rows(min_row=1,values_only=True):
        r+=1
        if r<8: continue
        c1=v[0] if len(v)>0 else None
        art=v[3] if len(v)>3 else None            # колонка D = Номенклатура.Артикул
        q =v[5] if len(v)>5 else None
        gr=v[6] if len(v)>6 else None
        rt=v[7] if len(v)>7 else None
        nt=v[8] if len(v)>8 else None
        if c1 is None and nt is None and q is None: continue
        rows.append([levels.get(r,0),(str(c1).strip() if c1 is not None else None),q,gr,rt,nt,(str(art).strip() if art else None)])
    wb.close()
    cache.parent.mkdir(exist_ok=True)
    cache.write_bytes(pickle.dumps(rows))
    print(f"  продажи ({path.name}): {len(rows)} строк за {round(time.time()-t0,1)}с")
    return rows

DOC=('заказ клиента','реализация','корректировка','возврат','поступление','перемещение','списание','оприходование','инвентаризация','отчет комиссионера')
_isdoc=lambda s:any(s.lower().startswith(k) for k in DOC)
_rx=re.compile(r'от (\d{2})\.(\d{2})\.(\d{4})')

def filter_year(rows, keep):
    """Оставляет только документы верхнего уровня, чья дата (год) входит в keep.
    Структурные заголовки (регион/субъект/город/компания) сохраняются; поддерево
    документа другого года (сам документ + вложенные строки) отбрасывается целиком.
    Позволяет взять 2024-2025 из полной истории и 2026 — из свежего файла без задвоения."""
    out=[]; skip=None; names={}
    for row in rows:
        lvl=row[0]; name=row[1]
        if skip is not None:
            if lvl>skip: continue       # внутри отбрасываемого поддерева
            skip=None                    # вышли на уровень документа-сиблинга/выше
        if name is None:
            names[lvl]=None
            for L in [x for x in names if x>lvl]: del names[L]
            out.append(row); continue
        parent=names.get(lvl-1)
        names[lvl]=name
        for L in [x for x in names if x>lvl]: del names[L]
        if _isdoc(name) and not _isdoc(parent or ''):   # только документ ВЕРХНЕГО уровня
            m=_rx.search(name)
            if m and int(m.group(3)) not in keep:
                skip=lvl; continue                        # отбрасываем документ и всё глубже
        out.append(row)
    return out

# округа Владивостока (остальные округа России -> Москва); СНГ -> ОПТ Москва; СЕТИ -> канал сети
VLAD_OKRUGA={'Дальневосточный ФО','Сибирский ФО (Восток)','Сибирский ФО (Запад)','Уральский ФО'}
SNG={'Беларусь','Казахстан','Кыргызстан'}
def _geo(stack):
    """По стеку определяем (регион/округ, филиал, канал)."""
    r0=stack.get(0)
    channel='СЕТИ' if r0=='СЕТИ' else 'ОПТ'
    if r0=='Россия':
        ok=stack.get(1)
        if ok=='Россия' or ok is None: return ('Не распределён','Не распределён','ОПТ')
        return (ok, 'Владивосток' if ok in VLAD_OKRUGA else 'Москва', 'ОПТ')
    if r0 in SNG:      return (r0,'Москва','ОПТ')          # СНГ ведёт ОПТ Москва
    if r0=='Москва':   return ('Москва','Москва','ОПТ')
    if r0=='СЕТИ':     return ('СЕТИ','Москва','СЕТИ')      # Москва-сети
    return ('Не распределён','Не распределён','ОПТ')

def parse_sales(rows, price):
    n=len(rows); stack={}
    art_cat={v[0]:v[3] for k,v in price.items() if k!='__art_price__' and isinstance(v,tuple) and v[0] and v[3]}
    company={'gross':0.0,'ret':0.0,'net':0.0,'qty':0}
    compYear=defaultdict(lambda:{'gross':0.0,'ret':0.0,'net':0.0,'qty':0})
    brandYear=defaultdict(lambda:defaultdict(float))
    byBrand=defaultdict(float)
    regSales=defaultdict(float); regClients=defaultdict(lambda:defaultdict(float))
    filSales=defaultdict(lambda:defaultdict(float))   # (год) -> филиал -> net; и канал
    chanSales=defaultdict(lambda:defaultdict(float))  # (год) -> канал(ОПТ/СЕТИ) -> net
    filChanYear=defaultdict(lambda:defaultdict(float))# год -> "филиал|канал" -> net
    clientRegion={}; clientFilial={}; clientChannel={}; clientCity={}; clientSubject={}
    client_net=defaultdict(float); client_qty=defaultdict(float)
    client_orders=defaultdict(int); client_last=defaultdict(str); client_orderhist=defaultdict(list)
    client_prod=defaultdict(lambda:defaultdict(lambda:[0.0,0.0]))
    comp=defaultdict(lambda:[0.0,0.0])
    yearMonth=defaultdict(float); ymGross=defaultdict(float); ymRet=defaultdict(float)
    posm_ct=[0,0.0]
    clientYear=defaultdict(lambda:defaultdict(float))
    tx=[]; tx_clients=[]; tx_cats=[]; tx_skus=[]; ci_m={}; ct_m={}; sk_m={}
    inwork=[]; maxdate=0; orderStack={}   # level -> (di,iso,num,osum, company, region, filial, channel)
    def ci(x):
        if x not in ci_m: ci_m[x]=len(tx_clients); tx_clients.append(x)
        return ci_m[x]
    def kt(c):
        c=c or 'Без категории'
        if c not in ct_m: ct_m[c]=len(tx_cats); tx_cats.append(c)
        return ct_m[c]
    def si(name,art=None):
        cat,art_p=cat_art(name,price)
        art=art or art_p                          # артикул из продаж; если нет — из прайса по имени
        if art and art_cat.get(art): cat=art_cat[art]   # категория по артикулу приоритетнее имени (KRON/ENKI матчатся по артикулу)
        _brand=brand_of(name,price)
        cat=group_cat(cat,_brand,name)                   # укрупняем KRON/ENKI (пены/герметики/клей/краски), fallback по названию
        key=art or ('n:'+nrm(name))
        if key not in sk_m: sk_m[key]=len(tx_skus); tx_skus.append([art or '—',name,kt(cat),_brand])
        return sk_m[key]
    def geo_for(companyname):
        reg,fil,chan=_geo(stack)
        low=(companyname or '').lower()
        if 'нордлогистик' in low: reg,fil='Сибирский ФО (Запад)','Владивосток'   # Новосибирск
        return reg,fil,chan
    for i in range(n):
        lvl,name,q,gr,rt,nt,art=rows[i]
        if name is None:
            stack[lvl]=None
            for L in [x for x in stack if x>lvl]: del stack[L]
            continue
        stack[lvl]=name
        for L in [x for x in stack if x>lvl]: del stack[L]
        for L in [x for x in orderStack if x>=lvl]: del orderStack[L]
        if lvl==0: continue
        low=name.lower()
        def cur_ord():
            ks=[k for k in orderStack if k<lvl]
            return orderStack[max(ks)] if ks else None
        parent=stack.get(lvl-1)
        if _isdoc(name):
            if not _isdoc(parent or ''):        # верхний документ компании (Заказ/Реализация/Комиссия…)
                comp_name=parent or stack.get(1) or 'Прочее'
                _cl=comp_name.lower()                         # маркетплейсы: канал СЕТИ, но не показывать как «компанию»
                if 'интернет решения' in _cl or '(озон)' in _cl or 'ozon' in _cl: comp_name='OZON (маркетплейс)'
                elif 'вайлдберриз' in _cl or 'wildberries' in _cl: comp_name='Wildberries (маркетплейс)'
                reg,fil,chan=geo_for(comp_name)
                if comp_name.strip().lower().startswith('частное лицо'):
                    comp_name=comp_name+' · '+reg   # не склеивать одноимённых по разным регионам
                m=_rx.search(name)
                di=int(f"{m.group(3)}{m.group(2)}{m.group(1)}") if m else 0
                iso=f"{m.group(3)}-{m.group(2)}-{m.group(1)}" if m else ''
                orderStack[lvl]=(di,iso,name.split('от')[0].strip(),nt or 0,comp_name,reg,fil,chan)
                clientRegion.setdefault(comp_name,reg); clientFilial.setdefault(comp_name,fil); clientChannel.setdefault(comp_name,chan)
                # город и субъект — из структуры отчёта продаж (ветка «Россия»: округ→субъект→город→клиент)
                _CL=lvl-1; _subj=''; _city=''
                if stack.get(0)=='Россия':
                    if _CL>=3: _subj=stack.get(2) or ''
                    if _CL>=4: _city=stack.get(3) or ''
                clientSubject.setdefault(comp_name,_subj); clientCity.setdefault(comp_name,_city)
                if di>maxdate: maxdate=di
                if iso and iso>client_last[comp_name]: client_last[comp_name]=iso
                if low.startswith('заказ клиента'):
                    client_orders[comp_name]+=1
                    if m:
                        client_orderhist[comp_name].append({'date':f"{m.group(1)}.{m.group(2)}.{m.group(3)}",
                            'num':name.split('от')[0].replace('Заказ клиента','').strip(),'sum':round(nt or 0),'qty':int(q or 0)})
                        if di//10000==CUR_YEAR:
                            shipped=0.0; j=i+1
                            while j<n and rows[j][0]>lvl and rows[j][1] is not None:
                                if rows[j][1].lower().startswith('реализация'): shipped+=rows[j][5] or 0
                                j+=1
                            ns=(nt or 0)-shipped
                            if ns>1: inwork.append({'client':comp_name,'order':name.split('от')[0].replace('Заказ клиента','').strip(),
                                'date':f"{m.group(1)}.{m.group(2)}.{m.group(3)}",'ordered':round(nt or 0),
                                'shipped':round(shipped),'not_shipped':round(ns)})
            else:
                # вложенный документ (Реализация под Заказом): датируем продажи ПО РЕАЛИЗАЦИИ,
                # а клиента/регион/филиал/канал/номер наследуем от родительского заказа
                base=cur_ord()
                if base:
                    m2=_rx.search(name)
                    di2=int(f"{m2.group(3)}{m2.group(2)}{m2.group(1)}") if m2 else base[0]
                    iso2=f"{m2.group(3)}-{m2.group(2)}-{m2.group(1)}" if m2 else base[1]
                    orderStack[lvl]=(di2,iso2,base[2],base[3],base[4],base[5],base[6],base[7])
                    if di2>maxdate: maxdate=di2
                    if iso2 and iso2>client_last[base[4]]: client_last[base[4]]=iso2
            continue
        nxt=rows[i+1][0] if i+1<n else -1
        if nxt>lvl: continue
        if nt is None: continue
        co=cur_ord()
        if not co: continue
        if is_posm(name):                          # POSM/непродаваемое (каталоги, пакеты, буклеты) — исключаем из аналитики
            posm_ct[0]+=1; posm_ct[1]+=(nt or 0); continue
        di,iso,onum,osum,client,reg,fil,chan=co
        q=q or 0; net=nt; b=brand_of(name,price)
        company['net']+=net; company['gross']+=(gr or 0); company['ret']+=(rt or 0); company['qty']+=q
        byBrand[b]+=net; regSales[reg]+=net; regClients[reg][client]+=net
        client_net[client]+=net; client_qty[client]+=q
        client_prod[client][name][0]+=q; client_prod[client][name][1]+=net
        comp[name][0]+=q; comp[name][1]+=net
        if di>0:
            yr=di//10000; ym=str(di)[:6]; yearMonth[ym]+=net; ymGross[ym]+=(gr or 0); ymRet[ym]+=(rt or 0)
            cy=compYear[yr]; cy['net']+=net; cy['gross']+=(gr or 0); cy['ret']+=(rt or 0); cy['qty']+=q
            brandYear[yr][b]+=net; clientYear[client][yr]+=net
            filSales[yr][fil]+=net; chanSales[yr][chan]+=net; filChanYear[yr][fil+'|'+chan]+=net
            if net!=0: tx.append([di,ci(client),si(name,art),round(q,1),round(net)])
    inwork.sort(key=lambda x:-x['not_shipped'])
    print(f"  исключено POSM/непродаваемого: {posm_ct[0]} строк, нетто {posm_ct[1]:,.0f} ₽")
    return dict(company=company,byBrand=dict(byBrand),regSales=dict(regSales),
        regClients={k:dict(v) for k,v in regClients.items()},clientRegion=clientRegion,
        clientFilial=clientFilial,clientChannel=clientChannel,clientCity=clientCity,clientSubject=clientSubject,
        client_net=dict(client_net),client_qty=dict(client_qty),client_orders=dict(client_orders),
        client_last=dict(client_last),client_orderhist=dict(client_orderhist),
        client_prod={k:dict(v) for k,v in client_prod.items()},comp=dict(comp),
        yearMonth=dict(yearMonth),ymGross=dict(ymGross),ymRet=dict(ymRet),clientYear={k:dict(v) for k,v in clientYear.items()},
        compYear={y:dict(d) for y,d in compYear.items()},brandYear={y:dict(d) for y,d in brandYear.items()},
        filSales={y:dict(d) for y,d in filSales.items()},chanSales={y:dict(d) for y,d in chanSales.items()},
        filChanYear={y:dict(d) for y,d in filChanYear.items()},
        tx=tx,tx_clients=tx_clients,tx_cats=tx_cats,tx_skus=tx_skus,inwork=inwork,maxdate=maxdate)

# ---------------- ОСТАТКИ (новый файл по сериям) ----------------
STOCK_WAREHOUSE='Адресный Лакония ОПТ'   # склад для оборачиваемости/остатков (как в 1С). «Всего доступно» для допродажи считается по ВСЕМ складам (company_in_stock/wh)
def load_stock():
    import openpyxl
    if not F_STOCK.exists(): print("  [!] нет файла остатков"); return {}
    # уровни группировки: склад (level 0) → серия (level 1) → строки артикула
    ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    levels={}
    try:
        with zipfile.ZipFile(F_STOCK).open('xl/worksheets/sheet1.xml') as fp:
            for ev,el in ET.iterparse(fp,events=('end',)):
                if el.tag==ns+'row':
                    rr=int(el.get('r')); ol=el.get('outlineLevel'); levels[rr]=int(ol) if ol else 0; el.clear()
    except Exception as e:
        print("  [!] остатки: не прочитал уровни группировки:",e)
    def wh_bucket(nm):
        n=(nm or '').lower()
        if 'лакония опт' in n: return 'Лакония ОПТ'
        if 'лакония сети' in n: return 'Лакония СЕТИ'
        if 'владивосток' in n or 'янковск' in n or 'мангут' in n: return 'Владивосток'
        if 'москва' in n or 'лист ложистик' in n or 'гидд' in n: return 'Москва'
        if 'новосибирск' in n: return 'Новосибирск'
        if 'лакония' in n: return 'Лакония прочее'
        return 'Прочее'
    wb=openpyxl.load_workbook(F_STOCK,read_only=True,data_only=True); ws=wb['Лист_1']
    num=lambda x:(x if isinstance(x,(int,float)) else 0)
    allrows=list(ws.iter_rows(min_row=1,values_only=True)); wb.close()
    # --- колонки определяем по заголовку (устойчиво к смене раскладки выгрузки 1С) ---
    COL={'art':0,'name':3,'unit':6,'in':7,'ship':8,'res':9,'avail':10}   # дефолт = формат «Остатки 08.10»
    hdr=next((rr for rr in allrows[:16] if any(isinstance(c,str) and 'в наличии' in c.lower() for c in rr)),None)
    if hdr:
        for i,c in enumerate(hdr):
            if not isinstance(c,str): continue
            t=c.lower().strip()
            if t.startswith('артикул'): COL['art']=i
            elif 'номенклат' in t: COL['name']=i
            elif 'ед. изм' in t: COL['unit']=i
            elif 'в наличии' in t: COL['in']=i
            elif 'отгружа' in t: COL['ship']=i
            elif 'в резерв' in t: COL['res']=i
            elif 'доступно' in t: COL['avail']=i
    # opt = по продажному складу (для оборачиваемости); comp/wh = по ВСЕМ складам (для «Всего доступно»)
    agg=defaultdict(lambda:{'name':'','in_stock':0.0,'shipping':0.0,'reserved':0.0,'available':0.0,'incoming':0.0,
                            'company':0.0,'wh':defaultdict(float)})
    cur_sklad=None
    for r,v in enumerate(allrows,1):
        g=lambda i:(v[i] if len(v)>i else None)
        c0=g(COL['art']); nmc=g(COL['name']); unc=g(COL['unit'])
        isart=bool(c0 and nmc and unc)                     # строка артикула: код + наименование + ед.изм
        if c0 and not isart:                               # заголовок (склад/серия)
            if levels.get(r,0)==0: cur_sklad=str(c0).strip()
            continue
        if isart:
            a=str(c0).strip()
            if a.lower()=='артикул' or str(nmc).strip().lower()=='номенклатура': continue  # строка-заголовок
            d=agg[a];
            if not d['name']: d['name']=str(nmc).strip()
            q7=num(g(COL['in']))
            d['company']+=q7                               # все склады
            if q7: d['wh'][wh_bucket(cur_sklad)]+=q7
            if cur_sklad==STOCK_WAREHOUSE:                 # только продажный склад
                d['in_stock']+=q7; d['shipping']+=num(g(COL['ship'])); d['reserved']+=num(g(COL['res']))
                d['available']+=num(g(COL['avail']))       # «в пути/в производстве» теперь из файла слежения (load_tracking)
    stock={a:{'name':d['name'],'in_stock':round(d['in_stock'],1),'shipping':round(d['shipping'],1),
              'reserved':round(d['reserved'],1),'available':round(d['available'],1),
              'incoming':round(d['incoming'],1),'company_in_stock':round(d['company'],1),
              'wh':{k:round(val,1) for k,val in d['wh'].items() if abs(val)>=0.5}} for a,d in agg.items()}
    print(f"  остатки: оборачиваемость по «{STOCK_WAREHOUSE}», «Всего доступно» по всем складам — {len(stock)} артикулов")
    return stock

# ---------------- СЛЕЖЕНИЕ / В ПУТИ (заказы поставщику, Китай) ----------------
def load_tracking(price):
    import openpyxl
    if not F_TRACK.exists(): print("  [!] нет файла слежения (в пути)"); return []
    ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'; lv={}
    try:
        with zipfile.ZipFile(F_TRACK).open('xl/worksheets/sheet1.xml') as fp:
            for ev,el in ET.iterparse(fp,events=('end',)):
                if el.tag==ns+'row':
                    rr=int(el.get('r')); ol=el.get('outlineLevel'); lv[rr]=int(ol) if ol else 0; el.clear()
    except Exception as e:
        print("  [!] слежение: уровни группировки:",e)
    def _d(v):
        if isinstance(v,str):
            m=re.match(r'(\d{2})\.(\d{2})\.(20\d{2})',v)
            if m: return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
        return None
    def _dest(wh):
        n=(wh or '').lower()
        if 'новосиб' in n or 'слк' in n: return 'Новосибирск'
        if 'сети' in n: return 'СЕТИ'
        if 'влад' in n or 'опт вл' in n or n.endswith(' вл') or 'янковск' in n or 'мангут' in n: return 'Владивосток'
        if 'мск' in n or 'москва' in n: return 'Москва'
        return 'Прочее'
    wb=openpyxl.load_workbook(F_TRACK,read_only=True,data_only=True)
    ws=wb[wb.sheetnames[0]]; allrows=list(ws.iter_rows(min_row=1,values_only=True)); wb.close()
    # --- колонки по заголовку (устойчиво к смене раскладки; артикул может отсутствовать) ---
    C={'name':8,'art':None,'qty':11,'port':3,'ktk':4,'eta':5}
    # строка-заголовок = та, где ячейка РОВНО «Номенклатура» (не путать с параметрами «…из списка номенклатуры»)
    hrow=next((rr for rr in allrows[:16] if any(isinstance(c,str) and c.strip().lower()=='номенклатура' for c in rr)),None)
    if hrow:
        for i,c in enumerate(hrow):
            if not isinstance(c,str): continue
            t=c.strip().lower()
            if t=='номенклатура': C['name']=i
            elif t.startswith('артикул'): C['art']=i
            elif 'порт' in t: C['port']=i
            elif 'ктк' in t: C['ktk']=i
            elif t.startswith('дата поступления') and 'ктк' not in t: C['eta']=i
    qrow=next((rr for rr in allrows[:16] if any(isinstance(c,str) and c.strip().lower()=='поступит' for c in rr)),None)
    if qrow:
        for i,c in enumerate(qrow):
            if isinstance(c,str) and c.strip().lower()=='поступит': C['qty']=i
    stack={}; agg={}
    for r,row in enumerate(allrows,1):
        if r<11: continue
        g=lambda i:(row[i] if len(row)>i else None)
        L=lv.get(r,0); a=g(0); name=g(C['name'])
        if a is not None and name is None:
            stack[L]=str(a).strip()
            for k in [x for x in list(stack) if x>L]: stack.pop(k,None)
            continue
        if name is None: continue
        doc=stack.get(1,'') or ''
        if not doc.lower().startswith('заказ поставщику'): continue   # только заказы поставщику (Китай)
        wh=stack.get(0,'') or ''; wl=wh.lower()
        status='В производстве' if wl.startswith('в производстве') else ('В пути' if wl.startswith('в пути') else 'Поступил')
        if status=='Поступил': continue   # уже на адресном складе — учитывается в остатках
        try: qty=float(g(C['qty'])) if g(C['qty']) not in (None,'') else 0
        except (TypeError,ValueError): qty=0
        if qty<=0: continue
        nm=str(name).strip()
        m=re.search(r'от (\d{2}\.\d{2}\.20\d{2})',doc); od=m.group(1) if m else ''
        port=_d(g(C['port'])); ktk=_d(g(C['ktk'])); eta=_d(g(C['eta']))
        _artf=g(C['art']) if C['art'] is not None else None  # артикул из файла слежения (если колонка есть)
        cat,art=cat_art(nm,price); br=brand_of(nm,price)
        art=(str(_artf).strip() if _artf not in (None,'') else '') or art   # приоритет артикула из слежения
        key=(art or nm, _dest(wh), status, eta or '', port or '', ktk or '', od)
        if key in agg: agg[key]['qty']+=qty
        else: agg[key]={'name':nm,'art':art or '','brand':br,'cat':group_cat(cat,br,nm),
            'qty':qty,'status':status,'dest':_dest(wh),'wh':wh,'order':od,'port':port,'ktk':ktk,'eta':eta}
    out=[dict(v,qty=round(v['qty'])) for v in agg.values()]
    out.sort(key=lambda x:(x['eta'] or '9999',x['status']))
    nprod=sum(1 for o in out if o['status']=='В производстве'); ntr=sum(1 for o in out if o['status']=='В пути')
    print(f"  в пути/в производстве: {len(out)} позиций заказов поставщику (в производстве {nprod}, в пути {ntr})")
    return out

# ---------------- СРОКИ ГОДНОСТИ (по сериям) ----------------
RAW_EXPIRY={}
def load_expiry(refiso):
    import openpyxl
    from datetime import datetime
    if not F_EXPIRY.exists(): print("  [!] нет файла сроков годности"); return {}
    ref=datetime(int(refiso[:4]),int(refiso[5:7]),int(refiso[8:]))
    wb=openpyxl.load_workbook(F_EXPIRY,read_only=True,data_only=True); ws=wb['Лист_1']
    def pdate(v):
        if isinstance(v,datetime): return v
        if isinstance(v,str):
            m=re.match(r'(\d{2})\.(\d{2})\.(\d{4})',v.strip())
            if m: return datetime(int(m.group(3)),int(m.group(2)),int(m.group(1)))
        return None
    batches=defaultdict(list); r=0
    for v in ws.iter_rows(min_row=1,values_only=True):
        r+=1
        if r<11: continue
        art=v[0] if len(v)>0 else None
        if not art: continue
        a=str(art).strip()
        if 'лакония' in a.lower() or 'склад' in a.lower(): continue   # строка склада
        d=pdate(v[9] if len(v)>9 else None)
        qty=v[11] if len(v)>11 else None
        if d is None: continue
        q=qty if isinstance(qty,(int,float)) else 0
        batches[a].append((d,q))
    wb.close()
    out={}
    horizon=90
    global RAW_EXPIRY
    RAW_EXPIRY={a:[[d.strftime('%Y-%m-%d'),q] for d,q in sorted((d,q) for d,q in bs if q>0)] for a,bs in batches.items()}
    for a,bs in batches.items():
        withq=sorted([(d,q) for d,q in bs if q>0])
        exp=round(sum(q for d,q in withq if d<ref))
        soon=round(sum(q for d,q in withq if ref<=d and (d-ref).days<=horizon))
        nearest=withq[0][0] if withq else None
        lst=[{'d':d.strftime('%d.%m.%Y'),'q':round(q),'st':('exp' if d<ref else('soon' if (d-ref).days<=horizon else 'ok'))} for d,q in withq]
        out[a]={'nearest':nearest.strftime('%d.%m.%Y') if nearest else None,
                'nearest_days':((nearest-ref).days if nearest else None),
                'expired':exp,'soon':soon,'list':lst}
    print(f"  сроки годности: {len(out)} артикулов с партиями")
    return out

# ---------------- CRM (ЦФО) ----------------
def _cnorm(s):
    s=str(s or '').lower().replace('ё','е')
    for ch in '«»"\'`(),.-–—/\\+№': s=s.replace(ch,' ')
    s=s.replace('новая компания',' ')
    LEG={'ооо','оао','зао','пао','ао','ип','тпк','тд','тк','нпо','нпф','гк','ooo'}
    return ' '.join(w for w in s.split() if w and w not in LEG)
def _clean(v):
    v=str(v or '').replace('&quot;','"').replace('&amp;','&').replace('&#160;',' ').replace('&nbsp;',' ')
    v=re.sub(r'\s+',' ',v).strip()
    if re.fullmatch(r'[\d,.]+[eE]\+\d+', v): return ''
    return v
def _crm_files():
    files=[]
    if CRM_DIR.exists():
        files += sorted(CRM_DIR.glob("*.xls"))
        f=CRM_DIR/"ЦФО.files"/"sheet001.htm"
        if f.exists(): files.append(f)
    old=ROOT/"ЦФО.files"/"sheet001.htm"
    if old.exists(): files.append(old)
    return files

def _parse_crm_file(path):
    """Разбор одного CRM-файла ПО НАЗВАНИЯМ колонок (разные округа = разные сдвиги)."""
    t=path.read_bytes().decode('utf-8','replace')
    rows=re.findall(r'<tr[^>]*>(.*?)</tr>',t,re.S|re.I)
    if len(rows)<2: return []
    cell=lambda r:[re.sub(r'<[^>]+>','',c) for c in re.findall(r'<t[hd][^>]*>(.*?)</t[hd]>',r,re.S|re.I)]
    hdr=[_clean(h) for h in cell(rows[0])]
    H=defaultdict(list)
    for i,h in enumerate(hdr): H[h].append(i)
    first=lambda name:(H[name][0] if H.get(name) else None)
    typ=H.get('Тип компании',[]); i_cat=(typ[-1] if typ else None); i_cat0=(typ[0] if typ else None)
    F={'name':first('Название компании'),'id':first('ID'),
       'inn':first('Реквизит (Россия): ИНН'),'kpp':first('Реквизит (Россия): КПП'),
       'ogrn':first('Реквизит (Россия): ОГРН'),
       'nameFull':first('Реквизит (Россия): Полное наименование организации'),
       'nameShort':first('Реквизит (Россия): Сокращенное наименование организации'),
       'director':first('Реквизит (Россия): Ген. директор'),'address':first('Реквизит (Россия): Адрес'),
       'phone1':first('Рабочий телефон'),'phone2':first('Мобильный телефон'),
       'email1':first('Рабочий e-mail'),'email2':first('Реквизит (Россия): E-Mail'),
       'site':first('Корпоративный сайт'),'responsible':first('Ответственный'),
       'employees':first('Кол-во сотрудников'),'contractNo':first('Номер договора'),
       'contractDate':first('Дата договора'),'bank':first('Банковский реквизит (Россия): Наименование банка'),
       'city':first('Город'),'subject':first('Субъект РФ'),
       'fam':first('Реквизит (Россия): Фамилия'),'im':first('Реквизит (Россия): Имя'),
       'ot':first('Реквизит (Россия): Отчество')}
    out=[]
    for r in rows[1:]:
        c=cell(r)
        gv=lambda i:(_clean(c[i]) if (i is not None and i<len(c)) else '')
        name=gv(F['name'])
        if not name: continue
        rec={k:gv(F[k]) for k in F}
        rec['name']=name; rec['category']=gv(i_cat) or gv(i_cat0)
        out.append(rec)
    return out

def load_crm(client_names, cinn=None):
    files=_crm_files()
    if not files: print("  CRM: файлы не найдены (папка Регионы)"); return {}
    recs=[]
    for p in files:
        ok=p.stem if p.suffix=='.xls' else 'ЦФО'
        rr=_parse_crm_file(p)
        for r in rr: r['_okrug']=ok
        recs+=rr
    # индекс: нормализованное имя -> запись (имя + юр.наименования + ФИО ИП)
    idx={}; inn_idx={}
    for rec in recs:
        fio=' '.join(x for x in [rec.get('fam',''),rec.get('im',''),rec.get('ot','')] if x).strip()
        for x in [rec.get('name',''),rec.get('nameShort',''),rec.get('nameFull',''),fio]:
            k=_cnorm(x)
            if k and k not in idx: idx[k]=rec
        _ri=re.sub(r'\D','',str(rec.get('inn','')))
        if _ri and _ri not in inn_idx: inn_idx[_ri]=rec
    tokidx=[(set(k.split()),rec) for k,rec in idx.items()]
    cid=lambda n:'c_'+re.sub(r'[^a-zа-я0-9]','',n.lower())[:20]
    cinn=cinn or {}; byInn=0
    out={}; matched=0
    for name in client_names:
        k=_cnorm(name); rec=idx.get(k)
        if not rec:
            ts=set(k.split())
            for kt2,rr in tokidx:
                if ts and (ts<=kt2 or kt2<=ts): rec=rr; break
        if not rec:                                   # по названию не нашли — пробуем по ИНН из справочника
            _ci=re.sub(r'\D','',str(cinn.get(nrm(name),'')))
            if _ci and inn_idx.get(_ci): rec=inn_idx[_ci]; byInn+=1
        if not rec: continue
        matched+=1
        phones=', '.join(x for x in [rec['phone1'],rec['phone2']] if x)
        emails=', '.join(dict.fromkeys(x for x in [rec['email1'],rec['email2']] if x))
        out[cid(name)]={'nameShort':rec['nameShort'] or rec['name'],'nameFull':rec['nameFull'],
            'inn':rec['inn'],'kpp':rec['kpp'],'ogrn':rec['ogrn'],'director':rec['director'],
            'address':rec['address'],'city':rec['city'],'subject':rec['subject'],'phone':phones,'email':emails,
            'site':rec['site'],'responsible':rec['responsible'],'category':rec['category'],
            'employees':rec['employees'],'contractNo':rec['contractNo'],'contractDate':rec['contractDate'],
            'bank':rec['bank'],'crmId':rec['id'],'okrug':rec.get('_okrug','')}
    print(f"  CRM: файлов {len(files)}, компаний {len(recs)}, сматчено {matched}/{len(client_names)} (из них по ИНН: {byInn})")
    return out

# ---------------- ПЛАН (25/26, помесячно) ----------------
_PMONTHS=['январь','февраль','март','апрель','май','июнь','июль','август','сентябрь','октябрь','ноябрь','декабрь']
def _plan_block(row, ci):
    """6 значений канала после названия месяца в позиции ci:
       [ВЛ ОПТ, МСК ОПТ, МСК СЕТИ] (KRONbuild+Enki) + [ВЛ, МСК, МСК СЕТИ] (HEADROCK)."""
    c=[x if isinstance(x,(int,float)) else 0 for x in row[ci+1:ci+7]]
    if len(c)<6: return None
    # c = [ВЛ ОПТ, МСК ОПТ, МСК СЕТИ (KRON+Enki)] + [ВЛ, МСК, МСК СЕТИ (HEADROCK)]
    return {'total':round(sum(c)),
            'vl':round(c[0]+c[3]),               # Владивосток (ОПТ) = ВЛ ОПТ + ВЛ
            'msk':round(c[1]+c[2]+c[4]+c[5]),     # Москва (всё)
            'mskopt':round(c[1]+c[4]),            # Москва ОПТ = МСК ОПТ + МСК(HR)
            'mskseti':round(c[2]+c[5]),           # Москва СЕТИ = МСК СЕТИ (обе марки)
            'kronenki':round(c[0]+c[1]+c[2]),     # KRONbuild+ENKI
            'headrock':round(c[3]+c[4]+c[5]),     # HEADROCK
            'cells':{'ОПТ Москва':{'kronenki':round(c[1]),'headrock':round(c[4])},
                     'ОПТ Владивосток':{'kronenki':round(c[0]),'headrock':round(c[3])},
                     'Сети':{'kronenki':round(c[2]),'headrock':round(c[5])}}}
def load_plan():
    import openpyxl
    if not F_PLAN.exists(): print("  [!] нет файла плана"); return {}
    wb=openpyxl.load_workbook(F_PLAN,data_only=True,read_only=True)
    plan={}
    # приоритет — сводный «Лист1» (там оба года и корректная разметка «Филиал»)
    sheets=(['Лист1'] if 'Лист1' in wb.sheetnames else [])+[s for s in wb.sheetnames if s.strip().isdigit()]
    def blank(): return {'months':{},'filials':{'Москва':{},'Владивосток':{}},
                         'brands':{'headrock':{},'kronenki':{}},
                         'buckets':{'ОПТ Москва':{},'ОПТ Владивосток':{},'Сети':{}},'cells':{},'annual':0}
    for sn in sheets:
        for v in wb[sn].iter_rows(values_only=True):
            for ci,cell in enumerate(v):
                if isinstance(cell,str) and cell.strip().lower() in _PMONTHS:
                    b=_plan_block(v,ci)
                    if not b: continue
                    # год: из соседнего блока справа тоже месяц? определяем по заголовку года не всегда есть.
                    # На «Лист1» левый блок=2025 (кол.D), правый=2026 (кол.K). На year-листах — имя листа.
                    if sn=='Лист1':
                        yr = 2025 if ci<7 else 2026
                    else:
                        try: yr=int(sn.strip())
                        except ValueError: continue
                    mi=_PMONTHS.index(cell.strip().lower())+1
                    p=plan.setdefault(yr,blank())
                    if mi in p['months']: continue
                    p['months'][mi]=b['total']; p['filials']['Москва'][mi]=b['msk']; p['filials']['Владивосток'][mi]=b['vl']
                    p['brands']['headrock'][mi]=b['headrock']; p['brands']['kronenki'][mi]=b['kronenki']
                    for _bk,_bg in b['cells'].items():
                        for _g,_v in _bg.items(): p['cells'].setdefault(_bk,{}).setdefault(_g,{})[mi]=_v
                    p['buckets']['ОПТ Москва'][mi]=b['mskopt']; p['buckets']['ОПТ Владивосток'][mi]=b['vl']; p['buckets']['Сети'][mi]=b['mskseti']
    for yr,p in plan.items():
        p['annual']=sum(p['months'].values())
        p['filialAnnual']={k:sum(m.values()) for k,m in p['filials'].items()}
        p['brandAnnual']={k:sum(m.values()) for k,m in p['brands'].items()}
        p['bucketAnnual']={k:sum(m.values()) for k,m in p['buckets'].items()}
    wb.close()
    yrs=", ".join(f"{y}: {plan[y]['annual']:,.0f}" for y in sorted(plan))
    print(f"  план: {yrs}")
    return plan

# ---------------- ABC ----------------
def abc_of(comp):
    items=sorted(((n,d[1]) for n,d in comp.items() if d[1]>0), key=lambda x:-x[1])
    tot=sum(v for _,v in items) or 1; cum=0.0; abc={}
    for n,v in items:
        cum+=v; sh=cum/tot
        abc[n]='A' if sh<=0.8 else('B' if sh<=0.95 else 'C')
    for n in comp:
        if n not in abc: abc[n]='C'
    return abc,tot

# ---------------- оборачиваемость (DOS) ----------------
def turnover(comp_recent_qty, stock, price):
    """DOS = В наличии / (продано за окно / дней окна). comp_recent_qty: name->qty за окно."""
    # свести продажи за окно к артикулу
    qty_by_art=defaultdict(float)
    for name,q in comp_recent_qty.items():
        _,art=cat_art(name,price)
        if art: qty_by_art[art]+=q
    out={}
    for art,s in stock.items():
        inst=s['in_stock'] or 0
        sold=qty_by_art.get(art,0)
        perday=sold/TURN_WINDOW_DAYS
        if inst<=0: continue
        dos=(inst/perday) if perday>0 else None   # None = мёртвый запас
        out[art]={'dos':(round(dos) if dos is not None else None),'in_stock':inst,
                  'sold_window':round(sold),'status':('dead' if dos is None else
                    ('short' if dos<30 else ('slow' if dos>180 else 'ok')))}
    return out

def _cid(n): return 'c_'+re.sub(r'[^a-zа-я0-9]','',str(n).lower())[:20]
def _daysago(iso, ref):
    if not iso: return None
    try: y,m,d=map(int,iso.split('-')); return (ref-date(y,m,d)).days
    except Exception: return None

def build_krs(S, stock, plan, crm, price):
    _exc={_cnorm(x) for x in EXCLUDE_CLIENTS_RAW}
    def excluded(nm): return (_cnorm(nm) in _exc) or (S['client_net'].get(nm,0)<=0)
    n_excluded=sum(1 for nm in S['client_net'] if excluded(nm))
    comp=S['comp']; abc,tot_rev = abc_of(comp)
    cy=S['compYear']; by=S['brandYear']
    CUR=CUR_YEAR if CUR_YEAR in cy else max(cy) if cy else CUR_YEAR
    PREV=CUR-1
    maxd=str(S['maxdate']); refiso=f"{maxd[:4]}-{maxd[4:6]}-{maxd[6:]}" if len(maxd)==8 else f"{CUR}-12-31"
    refdt=date(int(refiso[:4]),int(refiso[5:7]),int(refiso[8:]))
    expiry=load_expiry(refiso)
    # ---- дозаполнение «Не распределён» из CRM-округа + разрезы филиал/канал ----
    CRM_OKRUG={'УФО':('Уральский ФО','Владивосток'),'ДВФО':('Дальневосточный ФО','Владивосток'),
        'Западная Сибирь':('Сибирский ФО (Запад)','Владивосток'),'Восточная Сибирь':('Сибирский ФО (Восток)','Владивосток'),
        'ЦФО':('Центральный ФО','Москва'),'ЮФО':('Южный ФО','Москва'),'СЗФО':('Северо-Западный ФО','Москва'),
        'СКФО':('Северо-Кавказский ФО','Москва'),'ПФО':('Поволжский ФО','Москва'),
        'Москва и МО':('Москва','Москва'),'СНГ':('СНГ','Москва')}
    cReg=dict(S['clientRegion']); cFil=dict(S['clientFilial']); cChan=dict(S['clientChannel'])
    # CRM привязываем ПО ИНН (справочник _client_inn.xlsx); имя — только запасной матч
    _cinn,_kont=load_kontur()
    crmByInn={}
    for _rec in crm.values():
        _i=str((_rec or {}).get('inn') or '').strip()
        if _i and _i not in crmByInn: crmByInn[_i]=_rec
    def crm_of(nm):
        _i=_cinn.get(nrm(nm))
        if _i and _i in crmByInn: return crmByInn[_i]
        return crm.get(_cid(nm)) or {}
    filled=0
    for nm in list(cReg):
        if cReg[nm]=='Не распределён':
            rc=crm_of(nm); ok=(rc or {}).get('okrug','')
            if ok in CRM_OKRUG:
                cReg[nm],cFil[nm]=CRM_OKRUG[ok]; filled+=1
    # ручная привязка нераспределённых (регион/филиал); менеджер далее по региону
    _man_filled=0
    for nm in list(cReg):
        _mg=MANUAL_CLIENT_GEO.get(nrm(nm))
        if _mg: cReg[nm],cFil[nm]=_mg[0],_mg[1]; _man_filled+=1
    if _man_filled: print(f"  ручная привязка региона: {_man_filled}")
    SNG_REG={'Беларусь','Казахстан','Кыргызстан','СНГ'}
    def bucket(nm):
        if cChan.get(nm)=='СЕТИ': return 'Сети'
        if cReg.get(nm) in SNG_REG: return 'СНГ'
        f=cFil.get(nm)
        if f=='Владивосток': return 'ОПТ Владивосток'
        if f=='Москва': return 'ОПТ Москва'
        return 'Не распределён'
    print(f"  дозаполнено регионов из CRM: {filled}")
    # ---- менеджер по клиенту (закрепление округ→менеджер; сети — без менеджера) ----
    OKRUG_MGR={'Центральный ФО':'Сергей Сидоров','Северо-Западный ФО':'Сергей Сидоров',
        'Поволжский ФО':'Семён Комиссаров','Южный ФО':'Семён Комиссаров',
        'Дальневосточный ФО':'Александр Федотов','Сибирский ФО (Восток)':'Евгений Швайгерт',
        'Сибирский ФО (Запад)':'Дмитрий Жигалов','Уральский ФО':'Михаил Гнипель'}
    cMgr={}
    for nm in cReg:
        ch=cChan.get(nm); reg=cReg.get(nm)
        if ch=='СЕТИ': cMgr[nm]=None; continue          # у сетей менеджеров не берём
        if reg=='Москва':
            resp=(crm_of(nm) or {}).get('responsible','').strip()
            cMgr[nm]=resp if resp in ('Максим Чашников','Василий Димитрюк') else 'Максим Чашников'   # Москва без ответственного в CRM -> Чашников
        elif reg in ('Беларусь','Казахстан','Кыргызстан','СНГ'): cMgr[nm]='Сергей Сидоров'
        elif reg in OKRUG_MGR: cMgr[nm]=OKRUG_MGR[reg]
        else: cMgr[nm]='Не назначен'
    # ---- город/субъект по клиенту (ИЗ ПРОДАЖ: округ→субъект→город) ----
    cCity=dict(S.get('clientCity',{})); cSubj=dict(S.get('clientSubject',{}))
    for nm in list(cReg):
        _mg=MANUAL_CLIENT_GEO.get(nrm(nm))
        if _mg: cSubj[nm]=_mg[2]; cCity[nm]=_mg[3]
    # ---- конфликты: менеджер (список≠CRM), нет ИНН, нет CRM по ИНН ----
    conflicts=[]
    for nm in cReg:
        if cChan.get(nm)=='СЕТИ': continue
        inn=_cinn.get(nrm(nm)); rec=crm_of(nm); resp=(rec or {}).get('responsible','').strip(); mgr=cMgr.get(nm)
        if not inn: conflicts.append([nm,'нет ИНН в справочнике','',''])
        elif not rec: conflicts.append([nm,'нет записи CRM по ИНН',inn,''])
        if mgr and mgr!='Не назначен' and resp and resp!=mgr:
            conflicts.append([nm,'менеджер: список≠CRM',mgr,resp])
    try:
        import csv as _csv
        with open(ROOT/'work'/'_conflicts.csv','w',encoding='utf-8-sig',newline='') as _f:
            _w=_csv.writer(_f); _w.writerow(['Клиент','Тип','Список/ИНН','CRM'])
            for _r in conflicts: _w.writerow(_r)
    except Exception as _e: print('  [!] conflicts csv:',_e)
    from collections import Counter as _Cc
    print('  конфликты:',dict(_Cc(x[1] for x in conflicts)),'-> work/_conflicts.csv')
    # CRM для карточки — пересобрана по ИНН; ключ = _cid имени из продаж (как ждёт карточка)
    crmJoined={}
    for nm in S['client_net']:
        rec=crm_of(nm)
        if rec: crmJoined[_cid(nm)]=rec
    # ---- каталог + ABC (по продажам) ----
    cat_items=sorted(((n,d) for n,d in comp.items() if d[1]>0), key=lambda x:-x[1][1])
    catalog=[]; name2cat={}
    for n,d in cat_items:
        c,a=cat_art(n,price)
        name2cat[n]=len(catalog)
        catalog.append([a or '—', n, c or 'Без категории', abc.get(n,'C'), round(d[1])])
    A=[n for n in abc if abc[n]=='A']; B=[n for n in abc if abc[n]=='B']; C=[n for n in abc if abc[n]=='C']
    rev_by={'A':sum(comp[n][1] for n in A),'B':sum(comp[n][1] for n in B),'C':sum(comp[n][1] for n in C)}
    # покупатели по SKU (для резерва)
    buyers=defaultdict(int)
    for cl,prods in S['client_prod'].items():
        for n in prods:
            if prods[n][1]!=0: buyers[n]+=1
    nclients=len(S['client_prod'])
    a_sorted=sorted(A,key=lambda n:-comp[n][1])
    a_list=[{'article':(cat_art(n,price)[1] or '—'),'name':n,'rev':round(comp[n][1])} for n in a_sorted[:20]]
    reserve=[{'name':n,'nonbuyers':nclients-buyers[n],'rev':round(comp[n][1])}
             for n in sorted(A,key=lambda n:-((nclients-buyers[n])*comp[n][1]))[:15]]
    cat_total=defaultdict(int)
    for n in comp:
        c,_=cat_art(n,price)
        if c: cat_total[c]+=1
    # ---- оборачиваемость (DOS) из tx за окно ----
    cutoff=int((refdt.toordinal()-TURN_WINDOW_DAYS))
    def d2ord(di):
        s=str(di);
        try: return date(int(s[:4]),int(s[4:6]),int(s[6:])).toordinal()
        except Exception: return 0
    recent_qty_art=defaultdict(float)
    for t in S['tx']:
        if d2ord(t[0])>=cutoff:
            sk=S['tx_skus'][t[2]]; art=sk[0]
            if art and art!='—': recent_qty_art[art]+=t[3]
    art_price=price.get('__art_price__',{})
    turn={}
    for art,s in stock.items():
        inst=s['in_stock'] or 0
        if inst<=0: continue
        perday=recent_qty_art.get(art,0)/TURN_WINDOW_DAYS
        dos=(inst/perday) if perday>0 else None
        p1=art_price.get(art,0)
        frozen=round(inst*p1)
        ex=expiry.get(art,{})
        exp_q=ex.get('expired',0); soon_q=ex.get('soon',0)
        turn[art]={'dos':(round(dos) if dos is not None else None),'in_stock':inst,
                   'available':s['available'],'reserved':s.get('reserved',0),'shipping':s.get('shipping',0),'incoming':s.get('incoming',0),
                   'wh':s.get('wh',{}),'company_in_stock':s.get('company_in_stock',0),
                   'sold':round(recent_qty_art.get(art,0)),'frozen':frozen,'name':s['name'],
                   'exp_near':ex.get('nearest'),'exp_days':ex.get('nearest_days'),
                   'exp_expired':exp_q,'exp_soon':soon_q,
                   'exp_froz':round(exp_q*p1),'soon_froz':round(soon_q*p1),
                   'batches':ex.get('list',[]),
                   'status':('dead' if dos is None else('short' if dos<30 else('slow' if dos>180 else 'ok')))}
    # ---- overview (текущий год + план + АППГ) ----
    o_cur=cy.get(CUR,{'net':0,'gross':0,'ret':0,'qty':0}); o_prev=cy.get(PREV,{'net':0})
    plancur=plan.get(CUR,{}); planprev=plan.get(PREV,{})
    ym=S['yearMonth']
    def mfact(y,mi): return round(ym.get(f"{y}{mi:02d}",0)/1e6,3)
    months=[{'m':MN[mi-1],'fact':mfact(CUR,mi),
             'prev':mfact(PREV,mi),
             'plan':round(plancur.get('months',{}).get(mi,0)/1e6,3) if plancur else None} for mi in range(1,13)]
    clients_cur={cl:S['clientYear'].get(cl,{}).get(CUR,0) for cl in S['client_net'] if not excluded(cl)}
    clients_cur={cl:v for cl,v in clients_cur.items() if v}
    orders_cur=0
    for cl,hist in S['client_orderhist'].items():
        orders_cur+=sum(1 for h in hist if h['date'].endswith(str(CUR)))
    grow=lambda a,b: (round((a-b)/b*100,1) if b else None)
    cur_month=int(maxd[4:6]) if len(maxd)==8 else 12
    # граница данных: неполный последний месяц сравниваем с прошлым годом по ТЕМ ЖЕ датам, план — пропорционально дням
    import calendar
    _cut_d=int(maxd[6:8]) if len(maxd)==8 else 31
    _dim=calendar.monthrange(CUR,cur_month)[1]
    _partial=_cut_d<_dim
    _prev_cut=PREV*10000+cur_month*100+(_cut_d if _partial else 31)
    prev_same=sum(_t[4] for _t in S['tx'] if PREV*10000+101<=_t[0]<=_prev_cut)   # АППГ: тот же период прошлого года
    plan_ytd=sum(plancur.get('months',{}).get(mi,0)*((_cut_d/_dim) if (mi==cur_month and _partial) else 1) for mi in range(1,cur_month+1)) if plancur else None
    # АППГ для героя: клиенты/заказы/штуки за ТЕ ЖЕ месяцы прошлого года (из истории заказов)
    def _dm(ds):
        try: p=ds.split('.'); return (int(p[2]),int(p[1]))
        except: return (0,0)
    _oc=_op=_uc=_up=0; _cc=set(); _cp=set()
    for _cl,_hist in S['client_orderhist'].items():
        if excluded(_cl): continue
        for _h in _hist:
            _y,_mo=_dm(_h.get('date',''))
            if _mo<1 or _mo>cur_month: continue
            if _y==PREV and _mo==cur_month and _partial and int(_h.get('date','00.00.0000')[:2] or 0)>_cut_d: continue
            if _y==CUR: _oc+=1; _uc+=_h.get('qty',0); _cc.add(_cl)
            elif _y==PREV: _op+=1; _up+=_h.get('qty',0); _cp.add(_cl)
    yoy={'sales':grow(o_cur['net'],prev_same),'clients':grow(len(_cc),len(_cp)),'orders':grow(_oc,_op),'units':grow(_uc,_up)}
    # Обзор: категории/товары/регионы — ВСЁ за текущий год (CUR), чтобы совпадало с продажами года
    _catnet=defaultdict(float); _prodnet=defaultdict(float); _regnet=defaultdict(float)
    _cr=S.get('clientRegion',{})
    for _t in S['tx']:
        if _t[0]//10000!=CUR or _t[4]<=0: continue
        _sk=S['tx_skus'][_t[2]]
        _cn=S['tx_cats'][_sk[2]]
        if _cn and _cn!='Без категории': _catnet[_cn]+=_t[4]
        _prodnet[_sk[1]]+=_t[4]
        _rg=_cr.get(S['tx_clients'][_t[1]])
        if _rg and _rg!='Не распределён': _regnet[_rg]+=_t[4]
    _cats_sorted=sorted(_catnet.items(),key=lambda x:-x[1])[:8]
    _ctot=sum(v for _,v in _cats_sorted) or 1
    overview_categories=[{'name':c,'sales':round(v),'share':round(v/_ctot*100,1)} for c,v in _cats_sorted]
    overview_topproducts=[{'sku':(cat_art(n,price)[1] or '—'),'name':n,'abc':abc.get(n,'C'),'sales':round(v),
        'stock':(stock.get(cat_art(n,price)[1] or '',{}) or {}).get('in_stock',0)}
        for n,v in sorted(_prodnet.items(),key=lambda x:-x[1])[:10]]
    overview_regions=[{'name':r,'value':round(v)} for r,v in sorted(_regnet.items(),key=lambda x:-x[1])]
    overview={'sales':round(o_cur['net']),'grossSales':round(o_cur['gross']),'returns':round(o_cur['ret']),
        'units':int(o_cur['qty']),'returnsRate':round(-o_cur['ret']/o_cur['gross']*100,2) if o_cur['gross'] else 0,
        'plan':plancur.get('annual'),'planYTD':round(plan_ytd) if plan_ytd else None,
        'planDone':round(o_cur['net']/plan_ytd*100,1) if plan_ytd else None,
        'growth':grow(o_cur['net'],prev_same),
        'clients':len(clients_cur),'orders':orders_cur,'avgOrder':round(o_cur['net']/orders_cur) if orders_cur else 0,
        'yoy':yoy,
        'months':months,
        'branches':[{'id':'all','name':'Вся компания','sales':round(o_cur['net']),
            'plan':plancur.get('annual'),'growth':grow(o_cur['net'],o_prev.get('net',0)),
            'headrock':round(by.get(CUR,{}).get('HeadRock',0)),'kron':round(by.get(CUR,{}).get('KRONbuild',0)),
            'enki':round(by.get(CUR,{}).get('ENKI',0))}],
        'topClients':[{'name':n,'sales':round(v),'share':round(v/(o_cur['net'] or 1)*100,1)}
            for n,v in sorted(clients_cur.items(),key=lambda x:-x[1])[:10]],
        'topProducts':overview_topproducts,
        'regions':overview_regions,
        'categories':overview_categories}
    # ---- бренды (для salesDetail) ----
    brand_cur=by.get(CUR,{})
    salesDetail={'branches':{'moscow':{'name':'Вся компания','sales':round(o_cur['net']),'clients':len(clients_cur),
            'orders':orders_cur,'avgOrder':overview['avgOrder'],
            'brands':[{'id':'headrock','name':'HeadRock','fact':round(brand_cur.get('HeadRock',0))},
                      {'id':'kron','name':'KRONbuild','fact':round(brand_cur.get('KRONbuild',0))},
                      {'id':'enki','name':'ENKI','fact':round(brand_cur.get('ENKI',0))}]},
            'vlad':{'name':'Владивосток','sales':None,'missing':True}},
        'brands':{b.lower().replace('headrock','headrock').replace('kronbuild','kron'):{'name':b,
            'sales':round(brand_cur.get(b,0))} for b in ['HeadRock','KRONbuild','ENKI']},
        'categories':[{'name':c,'sales':round(sum(comp[n][1] for n in comp if cat_art(n,price)[0]==c)),
            'growth':None,'sku':cat_total[c]} for c in sorted(cat_total,key=lambda c:-cat_total[c])],
        'products':[]}
    salesDetail['brands']={'headrock':{'name':'HeadRock','sales':round(brand_cur.get('HeadRock',0)),'clients':len(clients_cur),'orders':orders_cur,'avgOrder':overview['avgOrder'],'brands':[]},
        'kron':{'name':'KRONbuild','sales':round(brand_cur.get('KRONbuild',0)),'clients':len(clients_cur),'orders':orders_cur,'avgOrder':overview['avgOrder'],'brands':[]},
        'enki':{'name':'ENKI','sales':round(brand_cur.get('ENKI',0)),'clients':len(clients_cur),'orders':orders_cur,'avgOrder':overview['avgOrder'],'brands':[]}}
    # ---- клиенты и карточки ----
    clients_list=[]; clientDetail={}; clientAssort={}; clientRegion=cReg
    A_set=set(A)
    for name in sorted(S['client_net'],key=lambda c:-S['clientYear'].get(c,{}).get(CUR,0)):
        if excluded(name): continue
        cid=_cid(name); prods=S['client_prod'].get(name,{})
        bought=set(n for n in prods if prods[n][1]!=0)
        rev_cur=S['clientYear'].get(name,{}).get(CUR,0)
        hist=S['client_orderhist'].get(name,[])
        hist_cur=[h for h in hist if h['date'].endswith(str(CUR))]
        ordn=len(hist_cur); last=S['client_last'].get(name,'')
        qty=round(sum(prods[n][0] for n in prods))
        # категории клиента
        catpen=defaultdict(lambda:[0,0.0])
        for n in bought:
            c,_=cat_art(n,price)
            if c: catpen[c][0]+=1; catpen[c][1]+=prods[n][1]
        cats=[]
        for c,(got,r) in sorted(catpen.items(),key=lambda x:-x[1][1]):
            totc=cat_total.get(c,got) or got; pen=round(got/totc*100) if totc else 0
            st='Хорошо' if pen>=50 else('Среднее' if pen>=25 else('Зона роста' if pen>0 else 'Не представлена'))
            cats.append({'name':c,'total':totc,'buy':got,'share':round(r/(sum(prods[x][1] for x in bought) or 1)*100,1),'status':st})
        cov={k:{'total':len(L),'buy':sum(1 for n in L if n in bought)} for k,L in (('A',A),('B',B),('C',C))}
        gaps=sorted([n for n in A if n not in bought],key=lambda n:-comp[n][1])[:12]
        rec=[{'sku':(cat_art(n,price)[1] or '—'),'name':n,'category':(cat_art(n,price)[0] or '—'),
              'brand':brand_of(n,price),'abc':'A','status':'Не покупает','action':'Добавить','effect':None,
              'orgSales':round(comp[n][1])} for n in gaps]
        days=_daysago(last,refdt); st='Активный' if (days is not None and days<=45) else('Риск' if (days or 999)>90 else 'Снижение')
        if rev_cur<=0: st='Риск'
        reg=clientRegion.get(name,'—')
        clients_list.append({'id':cid,'name':name,'manager':reg,'region':reg,'sales':round(rev_cur),
            'plan':None,'growth':grow(rev_cur,S['clientYear'].get(name,{}).get(PREV,0)),
            'brands':1,'categories':len(cats),'sku':len(bought),'last':last.replace('-','.') if last else '—',
            'potential':'Высокий' if len(gaps)>=100 else('Средний' if len(gaps)>=40 else 'Низкий'),'status':st,
            'avgCheck':round(rev_cur/ordn) if ordn else 0,'freq':ordn,'aGap':cov['A']['total']-cov['A']['buy'],
            'days':days})
        # months/categories/abc/recommended пересчитывает clientByPeriod из tx при открытии карточки —
        # в статике держим только лёгкую основу (иначе файл раздувается на мегабайты).
        _city=(cCity.get(name) or '').strip() or ((crm_of(name) or {}).get('city') or '').strip() or '—'
        clientDetail[cid]={'name':name,'manager':reg,'region':reg,'city':_city,'segment':'—','status':st,
            'sales':round(rev_cur),'units':qty,'orders':ordn,'avgOrder':round(rev_cur/ordn) if ordn else 0,
            'last':last.replace('-','.') if last else '—','growth':grow(rev_cur,S['clientYear'].get(name,{}).get(PREV,0)),
            'discount':None,'brands':[{'name':'HeadRock','sales':round(rev_cur),'share':100}],'months':[],
            'categories':[],'abc':cov,'recommended':[],
            'ordersHistory':[{'date':h['date'],'order':h['num'],'brand':'—','sum':h['sum'],'sku':h['qty'],
                'discount':None,'status':'Из отчёта'} for h in hist_cur[-12:]],'buySku':len(bought)}
        clientAssort[cid]={str(name2cat[n]):round(prods[n][1]) for n in bought if n in name2cat}
    # ---- регионы ----
    regionsData={}; _rc=defaultdict(lambda:defaultdict(float))
    for nm,net in S['client_net'].items():
        if excluded(nm): continue
        _rc[cReg.get(nm,'Не распределён')][nm]+=net
    for reg,cls in _rc.items():
        lst=sorted([{'name':n,'sales':round(v)} for n,v in cls.items() if v>0],key=lambda x:-x['sales'])
        regionsData[reg]={'sales':round(sum(cls.values())),'clients':len(lst),'list':lst}
    # факт по разрезам (ОПТ Москва / ОПТ Владивосток / СНГ / Сети) по годам + план
    factB=defaultdict(lambda:defaultdict(float))
    for t in S['tx']:
        nm=S['tx_clients'][t[1]]; factB[t[0]//10000][bucket(nm)]+=t[4]
    salesFilial={}; salesFilial3={}
    for y in sorted(set([CUR,PREV,CUR-2])):
        pl=plan.get(y,{}).get('bucketAnnual',{})
        f=lambda bk: round(factB.get(y,{}).get(bk,0))
        salesFilial[str(y)]={bk:{'fact':f(bk),'plan':(round(pl[bk]) if bk in pl else None)}
                             for bk in ('ОПТ Москва','ОПТ Владивосток','СНГ','Сети')}
        # 3 филиала: Москва = ОПТ Москва + СНГ; Владивосток; Сети
        salesFilial3[str(y)]={
            'Москва':{'fact':f('ОПТ Москва')+f('СНГ'),'plan':(round(pl['ОПТ Москва']) if 'ОПТ Москва' in pl else None)},
            'Владивосток':{'fact':f('ОПТ Владивосток'),'plan':(round(pl['ОПТ Владивосток']) if 'ОПТ Владивосток' in pl else None)},
            'Сети':{'fact':f('Сети'),'plan':(round(pl['Сети']) if 'Сети' in pl else None)}}
    foreign=[r for r in S['regSales'] if r in ('Беларусь','Казахстан','Кыргызстан')]
    regionsFull=[{'name':r,'sales':round(v)} for r,v in sorted(S['regSales'].items(),key=lambda x:-x[1])]
    # ---- остатки/дефицит ----
    stockOut={a:{'name':s['name'],'in_stock':s['in_stock'],'shipping':s['shipping'],
        'reserved_client':s['reserved'],'reserved_sale':0,'available':s['available'],
        'incoming':s['incoming'],'total_available':s['available']} for a,s in stock.items()}
    low=[{'article':a,'name':s['name'],'available':s['available'],'incoming':s['incoming'],'in_stock':s['in_stock']}
         for a,s in stock.items() if s['available']<=0]
    low.sort(key=lambda x:x['available']); low=low[:40]
    # ---- сборка ----
    # Контур.Фокус: ИНН по клиенту (имя из 1С — только ключ к справочнику; далее всё по ИНН); _cinn/_kont уже загружены выше
    clientInn={cn:_cinn[nrm(cn)] for cn in S['client_net'] if nrm(cn) in _cinn}
    print(f"  Контур: ИНН сопоставлен {len(clientInn)} клиентам из {len(S['client_net'])}")
    # транзакции исключённых клиентов (сотрудники из EXCLUDE_CLIENTS_RAW) не отдаём на фронт —
    # иначе страница «Менеджеры и клиенты» (считает из tx) снова их показывает
    _drop_ci={ix for ix,nm in enumerate(S['tx_clients']) if _cnorm(nm) in _exc}
    _tx_out=[t for t in S['tx'] if t[1] not in _drop_ci] if _drop_ci else S['tx']
    if _drop_ci: print(f"  tx: убрано строк исключённых клиентов {len(S['tx'])-len(_tx_out)}")
    KRS={
      'meta':{'prototype':True,'note':f'Пересобрано rebuild.py из актуальных Excel. Вся компания, {CUR} (АППГ {PREV}).',
        'sourceCoverage':{'period':f'{CUR} (АППГ {PREV})','scope':'Вся компания / все бренды',
          'actualNetSales':round(o_cur['net']),'actualUnits':int(o_cur['qty']),'actualReturns':round(o_cur['ret']),
          'missing':['маржинальность','себестоимость','скидки']}},
      'periods':[{'id':'ytd','label':f'{CUR} год','available':True}],
      'overview':overview,
      'abcSummary':{'totalRevenue':round(tot_rev),'counts':{'A':len(A),'B':len(B),'C':len(C)},'totalSku':len(catalog)},
      'salesDetail':salesDetail,'regions':{},'managers':{},
      'clients':clients_list,'clientDetail':clientDetail,'defaultClient':clients_list[0]['id'] if clients_list else '',
      'sourceLists':{'groups':regionsFull},
      'stock':stockOut,'ordersInWork':S['inwork'],'lowStock':low,
      'managersFull':{}, 'regionsFull':regionsFull,
      'abcDetail':{'counts':{'A':len(A),'B':len(B),'C':len(C)},'revenue':{k:round(v) for k,v in rev_by.items()},
        'revenue_share':{k:round(v/(tot_rev or 1)*100,1) for k,v in rev_by.items()},'total_rev':round(tot_rev),
        'total_sku':len(catalog),'nclients':nclients,'a_list':a_list,'reserve':reserve},
      'tx':_tx_out,'txClients':S['tx_clients'],'txCats':S['tx_cats'],'txSkus':S['tx_skus'],'catTotal':dict(cat_total),
      'ym':S['yearMonth'],'ymGross':S['ymGross'],'ymRet':S['ymRet'],'geoRef':load_geo(),'ruMap':load_rumap(),
      'regionsData':regionsData,'clientRegion':clientRegion,'foreignRegions':foreign or ['Беларусь'],
      'clientFilial':cFil,'clientChannel':cChan,'clientManager':cMgr,'clientCity':cCity,'clientSubject':cSubj,'salesFilial':salesFilial,'salesFilial3':salesFilial3,
      'refDate':refiso,'dataAsOf':refiso,'dataAsOfHuman':refiso[8:]+'.'+refiso[5:7]+'.'+refiso[:4],
      'catalog':catalog,'clientAssort':clientAssort,'crm':crmJoined,
      'clientInn':clientInn,'kontur':_kont,
      'turnover':turn,'plan':{str(y):plan[y] for y in plan},
    }
    print(f"  исключено служебных/нулевых: {n_excluded}")
    return KRS

def _inject_into(src_path, out_path, KRS):
    src=src_path.read_text(encoding='utf-8')
    key='window.KRS_DATA = '
    i=src.find(key)
    if i<0: raise SystemExit(f'слот window.KRS_DATA не найден в шаблоне {src_path.name}')
    j=i+len(key); _,rel=json.JSONDecoder().raw_decode(src[j:]); end=j+rel
    out=src[:j]+json.dumps(KRS,ensure_ascii=False)+src[end:]
    from selfcontain import make_self_contained
    out=make_self_contained(out,ROOT)   # картинки, шрифты и three.js внутри файла: открывается из любой папки и без интернета
    out_path.write_text(out,encoding='utf-8')
    print(f"  [OK] {out_path.name} пересобран ({len(out):,} байт)")

def export_refs(KRS, price, stock):
    """Справочники, нужные браузерному пересчёту (kairos_ingest.js): прайс, цены, остатки, партии сроков годности."""
    KRS['refs']={'price':[[k,v[0],v[1],v[2],v[3]] for k,v in price.items() if k!='__art_price__'],
                 'artPrice':price.get('__art_price__',{}),'stock':stock,'expiry':RAW_EXPIRY,
                 'stockWh':STOCK_WAREHOUSE,'turnWindow':TURN_WINDOW_DAYS}

def inject(KRS):
    try:
        import population_ref as _P
        KRS['pop']={'city':_P.POP_CITY,'region':_P.POP_REGION,'note':'оценка по Росстату (~2023-2024), округлённо; справочник population_ref.py'}
    except Exception as e:
        print('  [!] население не подключено:',e)
    _inject_into(TEMPLATE_SRC, OUT, KRS)
    proto_src=ROOT/"Kairos_proto_template.html"
    if proto_src.exists():
        proto_out=ROOT/"Kairos_dashboard_proto.html"
        _inject_into(proto_src, proto_out, KRS)
        # Боевой файл = актуальная версия proto. Держим final в синхроне,
        # чтобы он никогда не отдавал устаревшие данные (см. «всё сломалось»).
        try:
            import shutil as _sh
            _sh.copyfile(proto_out, TEMPLATE_SRC)
            print('  [=] Kairos_dashboard_final.html синхронизирован с proto')
        except Exception as _e:
            print('  [!] не удалось синхронизировать final.html:',_e)

if __name__=='__main__':
    print("Читаю источники…")
    price=load_price()
    # 2024-2025 — из полной истории; 2026 — из свежего файла (до актуальной даты), чтобы не задвоить 2026
    rows=filter_year(load_sales_rows(F_SALES, CACHE), {2024,2025})
    if F_SALES_CUR.exists():
        rows+=filter_year(load_sales_rows(F_SALES_CUR, CACHE_CUR), {2026})
        print(f"  слияние: 2024-2025 из истории + 2026 из свежего файла ({F_SALES_CUR.name})")
        # дневной файл («Продажи за ДД.ММ») не догружаем: он датируется по дате заказа и не сдвигает период —
        # для обновления продаж нужен ПОЛНЫЙ файл за 01.01.2026–<дата> (как «продажи 08.10»).
    else:
        rows+=filter_year(load_sales_rows(F_SALES, CACHE), {2026})
        print("  [!] свежего файла 2026 нет — 2026 взято из полной истории")
    print(f"  строк после слияния: {len(rows)}")
    S=parse_sales(rows,price)
    stock=load_stock()
    plan=load_plan()
    _cinn_main,_=load_kontur()
    crm=load_crm(list(S['client_net'].keys()),_cinn_main)
    print(f"  ИТОГО нетто: {S['company']['net']:,.0f} | клиентов: {len(S['client_net'])} | tx: {len(S['tx'])}")
    print(f"  бренды: "+", ".join(f"{k} {v:,.0f}" for k,v in sorted(S['byBrand'].items(),key=lambda x:-x[1])))
    print(f"  регионы: "+", ".join(f"{k} {v:,.0f}" for k,v in sorted(S['regSales'].items(),key=lambda x:-x[1])))
    pickle.dump({'S':S,'stock':stock,'plan':plan,'crm':crm,'price':price}, open(ROOT/"work"/"_parsed.pkl","wb"))
    print("Собираю модель…")
    KRS=build_krs(S,stock,plan,crm,price)
    KRS['incoming']=load_tracking(price)
    # таксономия товаров: артикул -> [группа, подгруппа]
    # подгруппа = категория прайса; если артикула нет в прайсе — берём категорию товара из модели (txCats),
    # чтобы продаваемые позиции вне прайса (напр. KRONbuild Smart) не падали в «Прочее».
    _priceCat={}
    for _k,_v in price.items():
        if _k=='__art_price__' or not isinstance(_v,tuple): continue
        if _v[0]: _priceCat[_v[0]]=_v[3]
    _tax={}; _tree={}
    _txc=KRS.get('txCats',[]); _nopr=0
    for _sk in KRS.get('txSkus',[]):
        _art=_sk[0] if _sk else None
        if not _art or _art=='—': continue
        _br=_sk[3] or 'HeadRock'; _catM=_txc[_sk[2]] if (_sk[2] is not None and _sk[2]<len(_txc)) else ''
        _raw=_priceCat.get(_art) or _catM or ''
        if _art not in _priceCat: _nopr+=1
        _g=group_of(_raw,_br)
        if _g=='Прочее':                       # нет категории — пробуем по названию (шуруп/дюбель→Крепёж и т.п.)
            _g2=group_of(_sk[1] or '',_br)
            if _g2!='Прочее': _g=_g2
        _s=_raw or '—'
        _tax[_art]=[_g,_s]; _tree.setdefault(_g,{})[_s]=1
    KRS['tax']=_tax
    KRS['groupTree']={_g:sorted(_tree[_g].keys()) for _g in _tree}
    print('  таксономия: групп %d, подгрупп %d, артикулов %d (вне прайса по категории: %d)'%(len(_tree),sum(len(v) for v in _tree.values()),len(_tax),_nopr))
    ov=KRS['overview']
    print(f"  overview {CUR_YEAR}: продажи {ov['sales']:,} ₽ | план YTD {ov['planYTD']:,} | "
          f"выполн {ov['planDone']}% | АППГ {ov['growth']}% (к тем же мес. {CUR_YEAR-1})")
    print(f"  клиентов {len(KRS['clients'])} | каталог {len(KRS['catalog'])} SKU | оборач. {len(KRS['turnover'])} арт.")
    export_refs(KRS,price,stock)
    inject(KRS)
