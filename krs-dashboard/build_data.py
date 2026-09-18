#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_data.py — считает ВСЕ данные для дашборда KRS/Кайрос из Excel-выгрузок.
Версия 2: с транзакциями (для фильтра по периоду), SKU-детализацией, остатками,
заказами в работе, ABC и категориями.

ЧТО ДЕЛАЕТ:
  1. Читает три Excel из data/:
     - Продажи по бизнес-регионам  (главный источник: транзакции до SKU с датами)
     - АВС по инструменту          (остатки склада: в наличии/резерв/доступно/поступит)
     - Прайс HR/KRON/ENKI          (справочник: бренд, категория, цена по артикулу)
  2. Считает по ВСЕМ клиентам: продажи, заказы, средний чек, ABC, матрицу категорий,
     проникновение, резерв A-позиций, историю заказов, помесячную динамику.
  3. Извлекает ТРАНЗАКЦИИ с датами (для фильтра период/месяц/диапазон на клиенте).
  4. Считает остатки и «заказы в работе» (заказ минус отгрузка).
  5. Пишет krs_data.json — весь массив данных для дашборда.
     Если рядом лежит KRS_template.html — вставляет данные в него → KRS_real_data.html

ОГРАНИЧЕНИЕ: выгрузка = Москва / HeadRock, 2026. Нет: плана, АППГ(2025), скидок,
маржи, продаж KRON/ENKI, Владивостока. Эти поля = null → в дашборде «нужен источник».
Сверка товар↔прайс идёт по НАЗВАНИЮ (~63%). Артикул в выгрузке продаж поднимет до 100%.

ЗАПУСК:  pip install openpyxl  &&  python build_data.py
"""
import openpyxl, json, re, sys
from collections import defaultdict
from pathlib import Path
from datetime import date

DATA=Path("data")
F_SALES=DATA/"Продажи_по_Бизнес_Регионам.xlsx"
F_ABC  =DATA/"Инструмент_клиенты_категории_А_по_АВС.xlsx"
F_PRICE=DATA/"Прайс.xlsx"
# CRM-выгрузка (Битрикс24, ЦФО) — HTML-таблица под видом .xls. Ищем sheet001.htm рядом.
CRM_CANDIDATES=[DATA/"ЦФО.files"/"sheet001.htm", Path("..")/"ЦФО.files"/"sheet001.htm",
                DATA/"sheet001.htm", Path("ЦФО.files")/"sheet001.htm"]
REF_DATE=date(2026,9,2)
MN={'01':'Янв','02':'Фев','03':'Мар','04':'Апр','05':'Май','06':'Июн','07':'Июл','08':'Авг','09':'Сен','10':'Окт','11':'Ноя','12':'Дек'}

def norm(s):
    s=str(s).lower().strip().replace('ё','е')
    s=re.sub(r'[«»"\'`(),]',' ',s); s=s.replace('х','x'); s=re.sub(r'\s+',' ',s)
    return s.strip()

# ---------- ПРАЙС ----------
def load_price():
    if not F_PRICE.exists():
        print(f"  [!] нет прайса {F_PRICE} — категории/бренды не посчитаются"); return {}, defaultdict(int)
    wb=openpyxl.load_workbook(F_PRICE,data_only=True,read_only=True); rows=[]
    if 'HeadRock' in wb.sheetnames:
        ws=wb['HeadRock']; cat=None
        for r in ws.iter_rows(min_row=9,values_only=True):
            art=r[1] if len(r)>1 else None; name=r[2] if len(r)>2 else None
            if art and (name is None or str(name).strip()==''): cat=str(art).strip(); continue
            if art and name: rows.append((str(art).strip(),str(name).strip(),'HeadRock',cat))
    if 'KRONbuild' in wb.sheetnames:
        ws=wb['KRONbuild']; cat=None
        for r in ws.iter_rows(min_row=10,values_only=True):
            art=r[0] if len(r)>0 else None; name=r[2] if len(r)>2 else None
            if (art is None or str(art).strip()=='') and name and str(name).strip(): cat=str(name).strip(); continue
            if art and name:
                brand='ENKI' if (cat and 'ENKI' in cat.upper()) else 'KRONbuild'
                rows.append((str(art).strip(),str(name).strip(),brand,cat))
    price={norm(n):(a,n,b,c) for a,n,b,c in rows}
    cat_total=defaultdict(int)
    for a,n,b,c in rows: cat_total[c or 'Без категории']+=1
    print(f"  прайс: {len(rows)} SKU"); return price, cat_total

# ---------- ПРОДАЖИ (главный файл) ----------
def load_sales(price):
    wb=openpyxl.load_workbook(F_SALES,data_only=True)
    ws=wb['Лист_1'] if 'Лист_1' in wb.sheetnames else wb.worksheets[0]
    def cat_art(n):
        m=price.get(norm(n)); return (m[3],m[0]) if m else (None,None)

    # Блок 1: менеджеры(lvl0)->клиенты(lvl1). qty4 sale6 ret8 total12. до 'Итого'
    managers=[]; cur=None; allclients={}; r=10
    while r<ws.max_row:
        name=ws.cell(row=r,column=1).value
        if name and str(name).strip()=='Итого': break
        if name:
            lvl=ws.row_dimensions[r].outlineLevel
            sale=ws.cell(row=r,column=6).value or 0; ret=ws.cell(row=r,column=8).value or 0; total=ws.cell(row=r,column=12).value or 0
            if lvl==0: cur={'m':str(name).strip(),'sale':sale,'ret':ret,'total':total,'clients':[]}; managers.append(cur)
            elif lvl==1 and cur: cur['clients'].append(str(name).strip()); allclients[str(name).strip()]={'manager':cur['m'],'ret':ret}
        r+=1
    clientset=set(allclients)
    itogo=None
    for rr in range(r-2,r+4):
        if ws.cell(row=rr,column=1).value and str(ws.cell(row=rr,column=1).value).strip()=='Итого': itogo=rr;break
    company={'sale':ws.cell(row=itogo,column=6).value or 0,'ret':ws.cell(row=itogo,column=8).value or 0,
             'total':ws.cell(row=itogo,column=12).value or 0,'qty':ws.cell(row=itogo,column=4).value or 0}

    # Блок 2: регионы->клиент->заказ->реализация->товар. границы:
    b2s=None;b2e=ws.max_row
    for rr in range(r,ws.max_row+1):
        v=ws.cell(row=rr,column=1).value
        if v and 'Партнер.Бизнес-регион' in str(v): b2s=rr+1
        elif b2s and v and str(v).strip()=='Итого': b2e=rr;break
    if not b2s:b2s=235
    levels={rr:ws.row_dimensions[rr].outlineLevel for rr in range(b2s,b2e)}
    rows_s=sorted(levels); idx={rr:i for i,rr in enumerate(rows_s)}
    def nm(rr):
        v=ws.cell(row=rr,column=1).value; return str(v).strip() if v else None
    client_row={}
    for rr in rows_s:
        s=nm(rr)
        if s in clientset and s not in client_row: client_row[s]=rr
    client_total={c:(ws.cell(row=client_row[c],column=18).value or 0) for c in client_row}

    DOC=('заказ клиента','реализация','возврат товаров','корректировка','поступление')
    is_doc=lambda s:any(s.lower().startswith(k) for k in DOC)
    rx=re.compile(r'от (\d{2})\.(\d{2})\.(\d{4})')

    # per-client products, orders, dates, transactions, orders-in-work
    client_prod=defaultdict(lambda:defaultdict(lambda:[0.0,0.0]))
    client_orders=defaultdict(int); client_last=defaultdict(str); client_orderhist=defaultdict(list)
    tx=[]; tx_clients=[]; tx_cats=[]; tx_skus=[]; ci_m={}; ct_m={}; sk_m={}
    def ci(x):
        if x not in ci_m: ci_m[x]=len(tx_clients); tx_clients.append(x)
        return ci_m[x]
    def kt(c):
        c=c or 'Без категории'
        if c not in ct_m: ct_m[c]=len(tx_cats); tx_cats.append(c)
        return ct_m[c]
    def si(name):
        cat,art=cat_art(name); key=art or ('n:'+norm(name))
        if key not in sk_m: sk_m[key]=len(tx_skus); tx_skus.append([art or '—',name,kt(cat)])
        return sk_m[key]
    inwork=[]
    for c,cr in client_row.items():
        clv=levels[cr]; i=idx[cr]+1; curdate=None
        while i<len(rows_s) and levels[rows_s[i]]>clv:
            rr=rows_s[i]; s=nm(rr); l=levels[rr]
            if s:
                if s.lower().startswith('заказ клиента'):
                    client_orders[c]+=1; m=rx.search(s)
                    order_sum=ws.cell(row=rr,column=18).value or 0
                    if m:
                        curdate=f"{m.group(3)}{m.group(2)}{m.group(1)}"; d=f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
                        if d>client_last[c]: client_last[c]=d
                        client_orderhist[c].append({'date':f"{m.group(1)}.{m.group(2)}.{m.group(3)}",
                            'num':s.split('от')[0].replace('Заказ клиента','').strip(),
                            'sum':round(order_sum),'qty':int(ws.cell(row=rr,column=11).value or 0)})
                    # orders in work: order minus child Реализация
                    shipped=0; j=i+1
                    while j<len(rows_s) and levels[rows_s[j]]>l:
                        cs=nm(rows_s[j])
                        if cs and cs.lower().startswith('реализация'): shipped+=ws.cell(row=rows_s[j],column=18).value or 0
                        j+=1
                    ns=order_sum-shipped
                    if ns>1 and m:
                        inwork.append({'client':c,'order':s.split('от')[0].replace('Заказ клиента','').strip(),
                            'date':f"{m.group(1)}.{m.group(2)}.{m.group(3)}",'ordered':round(order_sum),
                            'shipped':round(shipped),'not_shipped':round(ns)})
                elif not is_doc(s) and s not in clientset:
                    nxt=rows_s[i+1] if i+1<len(rows_s) else None
                    if (nxt is None) or (levels[nxt]<=l):  # leaf = product
                        q=ws.cell(row=rr,column=11).value or 0; v=ws.cell(row=rr,column=18).value or 0
                        client_prod[c][s][0]+=q; client_prod[c][s][1]+=v
                        if v!=0 and curdate: tx.append([int(curdate),ci(c),si(s),round(q,1),round(v)])
            i+=1
    inwork.sort(key=lambda x:-x['not_shipped'])
    return dict(managers=managers,allclients=allclients,client_total=client_total,
        client_prod={k:dict(v) for k,v in client_prod.items()},client_orders=dict(client_orders),
        client_last=dict(client_last),client_orderhist=dict(client_orderhist),company=company,
        tx=tx,tx_clients=tx_clients,tx_cats=tx_cats,tx_skus=tx_skus,inwork=inwork)

# ---------- ОСТАТКИ (файл ABC) ----------
def load_stock():
    if not F_ABC.exists(): print("  [!] нет файла ABC — остатки пропущены"); return {},[]
    wb=openpyxl.load_workbook(F_ABC,data_only=True); ws=wb['Лист_1']
    stock={}
    for r in range(13,701):
        art=ws.cell(row=r,column=4).value; name=ws.cell(row=r,column=6).value
        if not name: continue
        a=str(art).strip() if art else None
        row={'name':str(name).strip(),'in_stock':ws.cell(row=r,column=173).value or 0,
             'shipping':ws.cell(row=r,column=174).value or 0,'reserved_client':ws.cell(row=r,column=175).value or 0,
             'reserved_sale':ws.cell(row=r,column=176).value or 0,'available':ws.cell(row=r,column=177).value or 0,
             'incoming':ws.cell(row=r,column=178).value or 0,'total_available':ws.cell(row=r,column=179).value or 0}
        if a: stock[a]=row
    low=[{'article':a,'name':s['name'],'available':s['available'],'incoming':s['incoming'],'in_stock':s['in_stock']}
         for a,s in stock.items() if s['available']<=0]
    low.sort(key=lambda x:x['available'])
    print(f"  остатки: {len(stock)} позиций, дефицит: {len(low)}")
    return stock, low[:40]

# ---------- CRM (ЦФО, Битрикс24) ----------
def _cnorm(s):
    """Нормализация имени под матчинг: убираем юр.формы, пунктуацию."""
    s=str(s or '').lower().replace('ё','е')
    for ch in '«»"\'`(),.-–—/\\+№': s=s.replace(ch,' ')
    s=s.replace('новая компания',' ')
    LEG={'ооо','оао','зао','пао','ао','ип','тпк','тд','тк','нпо','нпф','гк','ooo'}
    return ' '.join(w for w in s.split() if w and w not in LEG)

def _clean(v):
    v=str(v or '').replace('&quot;','"').replace('&amp;','&').replace('&#160;',' ').replace('&nbsp;',' ')
    v=re.sub(r'\s+',' ',v).strip()
    # битые из Excel длинные числа (ОГРН, р/счёт) приходят как экспонента — отбрасываем
    if re.fullmatch(r'[\d,.]+[eE]\+\d+', v): return ''
    return v

def load_crm(client_names):
    """Читает CRM-htm, матчит по имени с клиентами дашборда, возвращает {cid: {...}}."""
    src=next((p for p in CRM_CANDIDATES if p.exists()), None)
    if not src:
        print("  CRM: файл ЦФО не найден — блок «Компания» будет пустым"); return {}
    t=src.read_bytes().decode('utf-8','replace')
    rows=re.findall(r'<tr[^>]*>(.*?)</tr>', t, re.S|re.I)
    def cells(r):
        cs=re.findall(r'<td[^>]*>(.*?)</td>', r, re.S|re.I)
        return [re.sub(r'<[^>]+>','',c) for c in cs]
    recs=[cells(r) for r in rows[1:]]
    recs=[c for c in recs if len(c)>158 and _clean(c[0]).isdigit() and _clean(c[2])]
    def g(c,i): return _clean(c[i]) if len(c)>i else ''
    # индекс: нормализованное имя -> запись CRM (имя + юр.наименования + ФИО ИП)
    idx={}
    for c in recs:
        cand=[g(c,2), g(c,157), g(c,158)]
        fio=' '.join([g(c,181),g(c,182),g(c,183)]).strip()
        if fio: cand.append(fio)
        for x in cand:
            k=_cnorm(x)
            if k and k not in idx: idx[k]=c
    tokidx=[(set(k.split()),c) for k,c in idx.items()]
    def cid(n): return 'c_'+re.sub(r'[^a-zа-я0-9]','',n.lower())[:20]
    out={}; matched=0
    for name in client_names:
        k=_cnorm(name); c=idx.get(k)
        if not c:  # запасной матч по подмножеству токенов
            ts=set(k.split())
            for kt,cc in tokidx:
                if ts and (ts<=kt or kt<=ts): c=cc; break
        if not c: continue
        matched+=1
        phones=', '.join(x for x in [g(c,11),g(c,12)] if x)
        emails=', '.join(dict.fromkeys(x for x in [g(c,25),g(c,166)] if x))
        out[cid(name)]={
            'nameShort':g(c,157) or g(c,2), 'nameFull':g(c,158),
            'inn':g(c,152),'kpp':g(c,160),'ogrn':g(c,159),'director':g(c,164),
            'address':g(c,167) or g(c,155),'city':g(c,56),'subject':g(c,83),
            'phone':phones,'email':emails,'site':g(c,18),'responsible':g(c,42),
            'category':g(c,68) or g(c,3),'employees':g(c,4),
            'contractNo':g(c,120),'contractDate':g(c,75),
            'bank':g(c,197),'crmId':g(c,0),
        }
    print(f"  CRM: компаний в выгрузке {len(recs)}, сматчено с клиентами {matched}/{len(client_names)}")
    return out

# ---------- СБОРКА ----------
def build():
    price,cat_total=load_price()
    S=load_sales(price)
    stock,low=load_stock()
    cp=S['client_prod']; ct=S['client_total']; ac=S['allclients']
    def cat_art(n):
        m=price.get(norm(n)); return (m[3],m[0]) if m else (None,None)
    # company ABC
    comp=defaultdict(lambda:[0.0,0.0])
    for c,prods in cp.items():
        for n,(q,v) in prods.items(): comp[n][0]+=q; comp[n][1]+=v
    pos=sorted([(n,d[1]) for n,d in comp.items() if d[1]>0],key=lambda x:-x[1])
    tp=sum(v for _,v in pos) or 1; cum=0; abc={}
    for n,v in pos:
        cum+=v; sh=cum/tp; abc[n]='A' if sh<=.8 else('B' if sh<=.95 else 'C')
    for n in comp: abc.setdefault(n,'C')
    A=[n for n in comp if abc[n]=='A']; B=[n for n in comp if abc[n]=='B']; C=[n for n in comp if abc[n]=='C']
    price_cat=defaultdict(set)
    # denominator per category from price
    def cid(n): return 'c_'+re.sub(r'[^a-zа-я0-9]','',n.lower())[:20]
    def daysago(ds):
        if not ds: return None
        y,m,d=map(int,ds.split('-')); return (REF_DATE-date(y,m,d)).days

    clients_list=[]; clientDetail={}
    for name in sorted(cp,key=lambda c:-ct[c]):
        prods=cp[name]; rev=ct[name]; bought=set(n for n in prods if prods[n][1]!=0)
        qty=sum(v[0] for v in prods.values()); o=S['client_orders'].get(name,0); last=S['client_last'].get(name,'')
        cov={k:{'total':len(L),'buy':sum(1 for n in L if n in bought)} for k,L in (('A',A),('B',B),('C',C))}
        catpen=defaultdict(lambda:[0,0.0])
        for n,v in prods.items():
            if v[1]==0:continue
            cat,_=cat_art(n)
            if cat: catpen[cat][0]+=1; catpen[cat][1]+=v[1]
        cats=[]
        for cat,(got,r) in sorted(catpen.items(),key=lambda x:-x[1][1]):
            tot=cat_total.get(cat,got) or got; pen=round(got/tot*100) if tot else 0
            st='Хорошо' if pen>=50 else('Среднее' if pen>=25 else('Зона роста' if pen>0 else 'Не представлена'))
            cats.append({'name':cat,'total':tot,'buy':got,'share':round(r/rev*100,1) if rev else 0,'status':st})
        gaps=sorted([n for n in A if n not in bought],key=lambda n:-comp[n][1])[:12]
        rec=[]
        for n in gaps:
            cat,art=cat_art(n)
            rec.append({'sku':art or '—','name':n,'category':cat or '—','brand':'HeadRock','abc':'A',
                        'status':'Не покупает','action':'Добавить','effect':None,'orgSales':round(comp[n][1])})
        cd=cid(name); st='Активный' if (daysago(last) is not None and daysago(last)<=45) else ('Риск' if (daysago(last) or 999)>90 else 'Снижение')
        if rev<=0: st='Риск'
        clients_list.append({'id':cd,'name':name,'manager':ac[name]['manager'],'region':'—','sales':round(rev),
            'plan':None,'growth':None,'brands':1,'categories':len(cats),'sku':len(bought),
            'last':last.replace('-','.') if last else '—',
            'potential':'Высокий' if len(gaps)>=100 else('Средний' if len(gaps)>=40 else 'Низкий'),'status':st})
        clientDetail[cd]={'name':name,'manager':ac[name]['manager'],'region':'—','city':'—','segment':'—','status':st,
            'sales':round(rev),'units':int(qty),'orders':o,'avgOrder':round(rev/o) if o else 0,
            'last':last.replace('-','.') if last else '—','growth':None,'discount':None,
            'brands':[{'name':'HeadRock','sales':round(rev),'share':100}],'months':[],'categories':cats,'abc':cov,
            'recommended':rec,'ordersHistory':[{'date':x['date'],'order':x['num'],'brand':'HeadRock','sum':x['sum'],
                'sku':x['qty'],'discount':None,'status':'Из отчёта'} for x in S['client_orderhist'].get(name,[])[-12:]],'buySku':len(bought)}
    # managers
    mgr={}
    for mm in S['managers']:
        cl=mm['clients']; orders=sum(S['client_orders'].get(c,0) for c in cl)
        mgr['m_'+re.sub(r'[^a-zа-я0-9]','',mm['m'].lower())[:16]]={'name':mm['m'],'sales':round(mm['total']),
            'ret':round(mm['ret']),'sale':round(mm['sale']),'clients':len(cl),'orders':orders,
            'avgOrder':round(mm['total']/orders) if orders else 0,
            'clientlist':[{'name':c,'sales':round(ct.get(c,0)),'orders':S['client_orders'].get(c,0)} for c in sorted(cl,key=lambda c:-ct.get(c,0))[:40]]}
    # ABC detail (revenue split, A-list, reserve) from company
    from collections import Counter
    rev_by={'A':sum(comp[n][1] for n in A),'B':sum(comp[n][1] for n in B),'C':sum(comp[n][1] for n in C)}
    tot=sum(rev_by.values()) or 1
    nclients=len(cp); buyers=defaultdict(int)
    for c,prods in cp.items():
        for n in set(nn for nn in prods if prods[nn][1]!=0): buyers[n]+=1
    a_sorted=sorted(A,key=lambda n:-comp[n][1])
    a_list=[{'article':cat_art(n)[1] or '—','name':n,'rev':round(comp[n][1])} for n in a_sorted[:20]]
    reserve=sorted(A,key=lambda n:-((nclients-buyers[n])*comp[n][1]))[:15]
    reserve=[{'name':n,'nonbuyers':nclients-buyers[n],'rev':round(comp[n][1])} for n in reserve]
    co=S['company']
    C0=defaultdict(lambda:[0,0.0])
    for n,d in comp.items():
        cat,_=cat_art(n)
        if cat: C0[cat][0]+=1; C0[cat][1]+=d[1]
    top_cat=sorted(C0.items(),key=lambda x:-x[1][1])[:8]; csum=sum(v[1] for _,v in top_cat) or 1
    KRS={
      'meta':{'prototype':True,'note':'Реальные данные: Москва / HeadRock, 2026.',
        'sourceCoverage':{'period':'2026','scope':'Москва / HeadRock','actualNetSales':round(co['total']),
          'actualUnits':int(co['qty']),'actualReturns':round(co['ret']),'missing':['планы','АППГ','скидки','маржинальность','KRONbuild','ENKI','Владивосток']}},
      'periods':[{'id':'ytd','label':'2026 год','available':True}],
      'overview':{'sales':round(co['total']),'grossSales':round(co['sale']),'returns':round(co['ret']),
        'units':int(co['qty']),'plan':None,'growth':None,'clients':len(clients_list),
        'orders':sum(S['client_orders'].values()),'avgOrder':round(co['total']/max(sum(S['client_orders'].values()),1)),
        'returnsRate':round(-co['ret']/co['sale']*100,2) if co['sale'] else 0,'months':[],
        'categories':[{'name':c,'sales':round(v[1]),'share':round(v[1]/csum*100,1)} for c,v in top_cat],
        'topClients':[{'name':x['name'],'sales':x['sales'],'share':round(x['sales']/co['total']*100,1)} for x in clients_list[:5]],
        'topProducts':[{'sku':cat_art(n)[1] or '—','name':n,'abc':abc[n],'sales':round(comp[n][1]),'stock':(stock.get(cat_art(n)[1],{}) or {}).get('in_stock',0)} for n,_ in sorted(comp.items(),key=lambda x:-x[1][1])[:10]]},
      'abcSummary':{'totalRevenue':round(tot),'counts':{'A':len(A),'B':len(B),'C':len(C)},'totalSku':len(A)+len(B)+len(C)},
      'abcDetail':{'counts':{'A':len(A),'B':len(B),'C':len(C)},'revenue':{k:round(v) for k,v in rev_by.items()},
        'revenue_share':{k:round(v/tot*100,1) for k,v in rev_by.items()},'total_rev':round(tot),
        'total_sku':len(A)+len(B)+len(C),'nclients':nclients,'a_list':a_list,'reserve':reserve},
      'salesDetail':{'branches':{},'brands':{},'categories':[{'name':c,'sales':round(v[1]),'growth':None,'sku':v[0]} for c,v in sorted(C0.items(),key=lambda x:-x[1][1])],'products':[]},
      'regions':{},'managers':{},'managersFull':mgr,
      'regionsFull':[{'name':n,'sales':v} for n,v in sorted([('Южный ФО',12828811),('Москва',15286375),('Поволжский ФО',12750812),('Центральный ФО',6667087),('Северо-Западный ФО',4410805),('Беларусь',5751611)],key=lambda x:-x[1])],
      'clients':clients_list,'clientDetail':clientDetail,'defaultClient':clients_list[0]['id'] if clients_list else '',
      'stock':stock,'lowStock':low,'ordersInWork':S['inwork'],
      'tx':S['tx'],'txClients':S['tx_clients'],'txCats':S['tx_cats'],'txSkus':S['tx_skus'],'catTotal':dict(cat_total),
      'crm':load_crm([c['name'] for c in clients_list]),
    }
    # monthly company dynamics from tx
    bm=defaultdict(float)
    for t in S['tx']:
        if 20260101<=t[0]<=20261231: bm[str(t[0])[:6]]+=t[4]
    KRS['overview']['months']=[{'m':MN[m],'fact':round(bm.get('2026'+m,0)/1e6,3),'plan':None,'prev':None} for m in ['01','02','03','04','05','06','07','08','09']]
    return KRS

def main():
    try: sys.stdout.reconfigure(encoding="utf-8")
    except Exception: pass
    if not F_SALES.exists():
        print(f"[X] нет файла продаж: {F_SALES}\n    положи выгрузки в data/ и проверь имена вверху скрипта"); sys.exit(1)
    print("Считаю..."); KRS=build()
    Path("krs_data.json").write_text(json.dumps(KRS,ensure_ascii=False),encoding="utf-8")
    print(f"\n[OK] krs_data.json записан")
    print(f"     клиентов: {len(KRS['clients'])}, транзакций: {len(KRS['tx'])}, SKU: {len(KRS['txSkus'])}")
    print(f"     продажи: {KRS['overview']['sales']:,} ₽, заказов: {KRS['overview']['orders']}, дефицит: {len(KRS['lowStock'])}")
    tmpl=Path("KRS_template.html")
    if tmpl.exists():
        html=tmpl.read_text(encoding="utf-8"); key='window.KRS_DATA = '
        i=html.find(key)+len(key); dec=json.JSONDecoder(); _,rel=dec.raw_decode(html[i:]); j=i+rel
        html=html[:i]+json.dumps(KRS,ensure_ascii=False)+html[j:]
        Path("KRS_real_data.html").write_text(html,encoding="utf-8")
        print("[OK] KRS_real_data.html пересобран из шаблона")
    else:
        print("     (шаблон KRS_template.html не найден — обновлён только krs_data.json)")

if __name__=="__main__": main()
